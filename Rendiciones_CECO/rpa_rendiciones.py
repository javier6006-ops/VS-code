# -*- coding: utf-8 -*-
"""
RPA RENDICIONES — Desglose por concepto de los respaldos adjuntos
=================================================================
Archivo único. Doble clic (o `python rpa_rendiciones.py`) abre una ventana para:
  1. Seleccionar el Excel "Presupuestos_Rendiciones_actualizado.xlsx".
  2. Elegir la hoja: "Detalle Adjuntos a Revisar" (32 auditados) o "Reembolsos Consolidado" (todas).
  3. Presionar "Iniciar": el RPA abre cada link de la columna Adjunto, descarga el archivo
     (PDF con texto o escaneado, fotos, Excel, Word, ZIP/RAR/7z anidados), lee los comprobantes
     y genera un Excel con el desglose por concepto y alertas de auditoría, junto al Excel de entrada.

Requisitos: Python 3.10+ (en Windows, el de python.org ya incluye la ventana Tkinter).
Las librerías faltantes se instalan solas la primera vez (requiere internet).
Opcionales:
  - Tesseract OCR (motor Local, para fotos y escaneados): https://github.com/UB-Mannheim/tesseract/wiki
    (en la instalación marcar idioma "Spanish").
  - 7-Zip (para abrir .rar): https://www.7-zip.org
  - API key de Anthropic (motor Claude, lectura muy superior de fotos/manuscritos).

Modo sin ventana:  python rpa_rendiciones.py --cli "ruta.xlsx" [--hoja ...] [--motor local|claude]
                   [--solo-movilizacion] [--limite N] [--reprocesar]
"""
from __future__ import annotations

# =====================================================================
# 0. DEPENDENCIAS (autoinstalación)
# =====================================================================
import importlib
import subprocess
import sys

_DEPS = {  # módulo -> paquete pip
    "openpyxl": "openpyxl", "pandas": "pandas", "xlrd": "xlrd", "requests": "requests",
    "pymupdf": "pymupdf", "PIL": "pillow", "pillow_heif": "pillow-heif", "pytesseract": "pytesseract",
    "docx": "python-docx", "rarfile": "rarfile", "py7zr": "py7zr", "anthropic": "anthropic",
}


def ensure_dependencies() -> None:
    missing = []
    for mod, pkg in _DEPS.items():
        try:
            importlib.import_module(mod)
        except ImportError:
            missing.append(pkg)
    if missing:
        print(f"Instalando librerías faltantes: {', '.join(missing)} ...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", *missing])
        importlib.invalidate_caches()


ensure_dependencies()

# =====================================================================
# IMPORTS
# =====================================================================
import argparse  # noqa: E402
import base64  # noqa: E402
import datetime as dt  # noqa: E402
import gzip  # noqa: E402
import hashlib  # noqa: E402
import io  # noqa: E402
import json  # noqa: E402
import logging  # noqa: E402
import os  # noqa: E402
import queue  # noqa: E402
import re  # noqa: E402
import shutil  # noqa: E402
import tempfile  # noqa: E402
import threading  # noqa: E402
import time  # noqa: E402
import unicodedata  # noqa: E402
import zipfile  # noqa: E402
from collections import defaultdict  # noqa: E402
from concurrent.futures import ThreadPoolExecutor, as_completed  # noqa: E402
from dataclasses import asdict, dataclass, field  # noqa: E402
from pathlib import Path  # noqa: E402
from urllib.parse import unquote, urlparse  # noqa: E402

import openpyxl  # noqa: E402
import requests  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from PIL import Image, ImageOps  # noqa: E402

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except ImportError:
    pass

log = logging.getLogger("rpa")

# =====================================================================
# 1. CONFIGURACIÓN (ajustable aquí mismo)
# =====================================================================
CONFIG = {
    # Estructura del Excel entregado: la hoja tiene una nota en la fila 1 y encabezados en la fila 2.
    # Se busca el encabezado automáticamente, así que funciona con ambas hojas.
    "hojas_soportadas": ["Detalle Adjuntos a Revisar", "Reembolsos Consolidado"],
    "hoja_auditoria_cabify": "Auditoria Cabify vs Rend",
    "estados_pagados": ["finalizado", "pagada"],      # filtro "Sólo pagadas/finalizadas"
    # Botón "CECOs operacionales" en la ventana / --cecos-operacionales en modo CLI
    "cecos_operacionales": ["220 - GERENCIA OPERACIONES", "221 - RECLUTAMIENTO Y SELECCIÓN", "230 - SANTIAGO",
                            "231 - SUPPLY CHAIN", "240 - ANTOFAGASTA", "250 - LA SERENA", "260 - VIÑA DEL MAR",
                            "265 - TALCA", "270 - CONCEPCIÓN", "280 - PUERTO MONTT", "290 - TRADE", "295 - RETAIL",
                            "298 - FACILITY", "320 - BUSINESS PROFESSIONAL"],
    "columnas": {  # nombres posibles de cada campo (se usa el primero que exista)
        "id":       ["ID"],
        "rendidor": ["Rendidor", "Nombre Creador"],
        "ceco":     ["CECO", "Indique el CECO"],
        "concepto": ["Concepto", "Indique cuenta contable"],
        "fecha":    ["Fecha"],
        "estado":   ["Estado", "Indique estado de la solicitud"],
        "monto":    ["Monto", "Indique monto a sustentar o reembolsar"],
        # En "Detalle Adjuntos a Revisar" el link es un hipervínculo con texto "Abrir archivo";
        # en "Reembolsos Consolidado" la URL viene como texto en esta columna larga.
        "adjunto":  ["Adjunto", "Adjuntar el archivo (en caso de necesitar subir más de un archivo, "
                                "se debe comprimir la información en un archivo .RAR)"],
    },
    "claude": {"modelo": "claude-sonnet-5",           # más barato: "claude-haiku-4-5-20251001"
               "max_tokens": 4000, "max_chars_texto": 30000, "lado_max_imagen_px": 1568},
    "local": {"idioma_ocr": "spa", "dpi_pdf_escaneado": 250,
              "tesseract_rutas": [r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                                  r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                                  os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe")]},
    "descarga": {"timeout_seg": 60, "reintentos": 3, "workers": 4},
    "descompresion": {"profundidad_max": 3, "max_archivos_por_adjunto": 200,
                      "rutas_7zip": [r"C:\Program Files\7-Zip\7z.exe", r"C:\Program Files (x86)\7-Zip\7z.exe"]},
    "documentos": {"max_paginas_pdf": 40, "min_chars_pdf_con_texto": 40},
    "auditoria": {"tolerancia_diferencia_pct": 5, "dias_max_antiguedad_doc": 90,
                  "conceptos_alerta": ["COMBUSTIBLE", "APPS_TRANSPORTE", "TAXI_COLECTIVO"]},
    # Taxonomía: el ORDEN define la prioridad (primer concepto que calza gana). Palabras sin tildes.
    "conceptos": {
        "APPS_TRANSPORTE": ["cabify", "uber", "didi", "indriver", "indrive", "beat"],
        "COMBUSTIBLE": ["copec", "shell", "petrobras", "aramco", "enex", "terpel", "esso", "bencina", "gasolina",
                        "combustible", "diesel", "petroleo", "kerosene", "gas 93", "gas 95", "gas 97",
                        "sin plomo", "lts", "litros"],
        "ESTACIONAMIENTO": ["estacionamiento", "parking", "parquimetro", "saba", "central parking"],
        "PEAJE_TAG": ["peaje", "autopista", "costanera", "vespucio", "autopase", "tag", "ruta 68", "ruta 5",
                      "vias chile", "globalvia", "pase diario"],
        "TAXI_COLECTIVO": ["taxi", "colectivo", "radiotaxi", "transfer"],
        "TRANSPORTE_PUBLICO": ["metro", "bip", "red movilidad", "transantiago", "tren", "efe", "merval", "biotren"],
        "BUS_INTERURBANO": ["turbus", "tur bus", "pullman", "buses", "jac", "condor", "eme bus", "pasaje bus"],
        "PASAJE_AEREO": ["latam", "sky airline", "jetsmart", "pasaje aereo", "vuelo"],
        "ALOJAMIENTO": ["hotel", "hostal", "alojamiento", "airbnb", "cabana", "residencial", "booking"],
        "ALIMENTACION": ["restaurant", "restaurante", "almuerzo", "cena", "desayuno", "colacion", "comida", "cafe",
                         "cafeteria", "sandwich", "menu", "pizza", "sushi", "mcdonald", "burger", "kfc",
                         "starbucks", "juan maestro", "doggis", "bebida", "agua mineral"],
        "UTILES_OFICINA": ["libreria", "resma", "lapiz", "cuaderno", "toner", "tinta", "carpeta", "papel",
                           "corchetera"],
        "NOTARIA": ["notaria", "notario", "legalizacion", "conservador"],
        "COURIER": ["chilexpress", "starken", "correos de chile", "bluexpress", "courier", "envio"],
        "FERRETERIA_MANTENCION": ["sodimac", "easy", "ferreteria", "construmart", "herramienta"],
        "SUPERMERCADO": ["lider", "jumbo", "unimarc", "santa isabel", "tottus", "supermercado", "acuenta"],
        "FARMACIA": ["farmacia", "cruz verde", "salcobrand", "ahumada"],
        "OTROS": [],
    },
}


# =====================================================================
# 2. UTILIDADES (texto, montos CLP, fechas, RUT, folio)
# =====================================================================
def normalize(s: object) -> str:
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s).strip().lower()


_AMOUNT_RE = re.compile(r"\$?\s?(\d{1,3}(?:[.,]\d{3})+(?:,\d{1,2})?|\d{3,9}(?:,\d{1,2})?)")


def parse_amount(raw: object) -> float | None:
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip().replace("$", "").replace(" ", "").replace("CLP", "")
    if not s:
        return None
    if "." in s and "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        parts = s.split(",")
        s = s.replace(",", "") if all(len(p) == 3 for p in parts[1:]) else s.replace(",", ".")
    elif "." in s and all(len(p) == 3 for p in s.split(".")[1:]):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def find_amounts(text: str, min_v: float = 100, max_v: float = 50_000_000) -> list[float]:
    out = []
    for m in _AMOUNT_RE.finditer(text):
        v = parse_amount(m.group(1))
        if v is not None and min_v <= v <= max_v:
            out.append(v)
    return out


_DATE_PATTERNS = [(re.compile(r"\b(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\b"), "dmy"),
                  (re.compile(r"\b(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})\b"), "ymd"),
                  (re.compile(r"\b(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2})\b"), "dmy2")]


def find_date(text: str) -> dt.date | None:
    for rx, kind in _DATE_PATTERNS:
        for m in rx.finditer(text):
            a, b, c = (int(x) for x in m.groups())
            try:
                return dt.date(c, b, a) if kind == "dmy" else dt.date(a, b, c) if kind == "ymd" else dt.date(2000 + c, b, a)
            except ValueError:
                continue
    return None


def parse_date(raw: object) -> dt.date | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, dt.datetime):
        return raw.date()
    if isinstance(raw, dt.date):
        return raw
    s = str(raw).strip()
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        return find_date(s)


_RUT_RE = re.compile(r"\b(\d{1,2}\.?\d{3}\.?\d{3}\s?-\s?[\dkK])\b")
_FOLIO_RE = re.compile(r"(?:folio|n[°ºo]\.?|nro\.?|(?:boleta|factura)\s+electronica\s+n)\s*[:#]?\s*(\d{3,12})", re.I)


def find_rut(text: str) -> str | None:
    m = _RUT_RE.search(text)
    return m.group(1).replace(".", "").replace(" ", "").upper() if m else None


def find_folio(text: str) -> str | None:
    m = _FOLIO_RE.search(text) or _FOLIO_RE.search(normalize(text))   # "N°" se pierde al normalizar
    return m.group(1) if m else None


def first_existing(paths: list[str]) -> str | None:
    return next((p for p in paths if p and Path(p).exists()), None)


SETTINGS_FILE = Path(__file__).resolve().parent / "rpa_rendiciones_settings.json"


def load_settings() -> dict:
    if SETTINGS_FILE.exists():
        try:
            return json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            pass
    return {}


def save_setting(key: str, value: str) -> None:
    data = load_settings()
    data[key] = value
    try:
        SETTINGS_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass


def _glob_search(exe_name: str, roots: list[str], max_dirs: int = 6000) -> str | None:
    """Busca exe_name bajo cada raíz, hasta 3 niveles de profundidad (evita escanear todo el disco)."""
    seen = 0
    for root in roots:
        base = Path(root)
        if not base.exists():
            continue
        for pattern in (exe_name, f"*/{exe_name}", f"*/*/{exe_name}"):
            for hit in base.glob(pattern):
                seen += 1
                if hit.is_file():
                    return str(hit)
                if seen > max_dirs:
                    return None
    return None


def find_tesseract() -> str | None:
    override = load_settings().get("tesseract_cmd")
    if override and Path(override).exists():
        return override
    found = first_existing(CONFIG["local"]["tesseract_rutas"]) or shutil.which("tesseract")
    if found:
        return found
    return _glob_search("tesseract.exe", [r"C:\Program Files", r"C:\Program Files (x86)",
                                          os.path.expandvars(r"%LOCALAPPDATA%\Programs"),
                                          os.path.expandvars(r"%USERPROFILE%")])


def find_7zip() -> str | None:
    override = load_settings().get("sevenzip_cmd")
    if override and Path(override).exists():
        return override
    found = first_existing(CONFIG["descompresion"]["rutas_7zip"]) or shutil.which("7z") or shutil.which("unrar")
    if found:
        return found
    return _glob_search("7z.exe", [r"C:\Program Files", r"C:\Program Files (x86)"])         or _glob_search("UnRAR.exe", [r"C:\Program Files", r"C:\Program Files (x86)",
                                      os.path.expandvars(r"%USERPROFILE%")])


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


# =====================================================================
# 3. LECTURA DEL EXCEL ENTREGADO
# =====================================================================
@dataclass
class Rendicion:
    id: str
    rendidor: str
    ceco: str
    concepto_declarado: str
    fecha: object
    monto_declarado: float | None
    adjunto: str | None
    usa_cabify: str = "NO"
    estado: str = ""


def _find_header(ws, required: list[str], max_scan: int = 10) -> tuple[int, dict[str, int]]:
    for r in range(1, max_scan + 1):
        values = [normalize(c.value) for c in ws[r]]
        if all(normalize(req) in values for req in required):
            return r, {v: i for i, v in enumerate(values) if v}
    raise ValueError(f"No encontré encabezado con {required} en la hoja '{ws.title}'")


def _cell_link(cell) -> str | None:
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target
    v = cell.value
    if isinstance(v, str):
        v = v.strip()
        if v.lower().startswith(("http://", "https://")) or (len(v) < 260 and Path(v).exists()):
            return v
    return None


def list_sheets(xlsx: Path) -> list[str]:
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    names = wb.sheetnames
    wb.close()
    return [s for s in CONFIG["hojas_soportadas"] if s in names] or names


def list_cecos(xlsx: Path, hoja: str) -> list[str]:
    """CECOs distintos de la hoja (para el selector de la ventana)."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        ws = wb[hoja]
        hr, hmap = _find_header(ws, ["ID"])
        col = next((hmap[normalize(c)] for c in CONFIG["columnas"]["ceco"] if normalize(c) in hmap), None)
        if col is None:
            return []
        return sorted({str(r[col]).strip() for r in ws.iter_rows(min_row=hr + 1, values_only=True)
                       if col < len(r) and r[col]})
    finally:
        wb.close()


def read_rendiciones(xlsx: Path, hoja: str, solo_movilizacion=False, rendidor=None, limite=None,
                     cecos: list[str] | None = None, solo_finalizadas: bool = False) -> list[Rendicion]:
    log.info("Leyendo '%s' (hoja '%s')...", xlsx.name, hoja)
    wb = openpyxl.load_workbook(xlsx, data_only=True)       # modo normal: necesario para leer hipervínculos
    if hoja not in wb.sheetnames:
        raise ValueError(f"La hoja '{hoja}' no existe. Hojas disponibles: {wb.sheetnames}")
    ws = wb[hoja]
    hr, hmap = _find_header(ws, ["ID"])
    idx = {}
    for k, cands in CONFIG["columnas"].items():
        idx[k] = next((hmap[normalize(c)] for c in cands if normalize(c) in hmap), None)
    missing = [k for k in ("id", "rendidor", "adjunto") if idx[k] is None]
    if missing:
        raise ValueError(f"En la hoja '{hoja}' faltan las columnas {missing}.")

    # usuarios Cabify desde la hoja de auditoría (si existe)
    cabify = set()
    sh = CONFIG["hoja_auditoria_cabify"]
    if sh in wb.sheetnames:
        try:
            ahr, amap = _find_header(wb[sh], ["¿Coincide en Cabify?"])
            ncol = next((i for k, i in amap.items() if k.startswith("nombre (origen: talana)")
                         or k.startswith("nombre creador")), 0)
            fcol = amap[normalize("¿Coincide en Cabify?")]
            for row in wb[sh].iter_rows(min_row=ahr + 1):
                if normalize(row[fcol].value) == "si" and row[ncol].value:
                    cabify.add(normalize(row[ncol].value))
        except ValueError:
            pass

    filtro = normalize(rendidor) if rendidor else None
    cecos_n = {normalize(c) for c in cecos} if cecos else None
    pagados = {normalize(e) for e in CONFIG["estados_pagados"]}
    if solo_finalizadas and idx.get("estado") is None:
        log.warning("La hoja no tiene columna de estado: no se puede filtrar por pagadas/finalizadas.")
    out: list[Rendicion] = []
    for row in ws.iter_rows(min_row=hr + 1):
        def get(k):
            return row[idx[k]].value if idx[k] is not None and idx[k] < len(row) else None
        rid = get("id")
        if rid in (None, ""):
            continue
        nombre = str(get("rendidor") or "").strip()
        concepto = str(get("concepto") or "").strip()
        if solo_movilizacion and "movilizacion" not in normalize(concepto):
            continue
        if filtro and filtro not in normalize(nombre):
            continue
        ceco = str(get("ceco") or "").strip()
        if cecos_n is not None and normalize(ceco) not in cecos_n:
            continue
        estado = str(get("estado") or "").strip()
        if solo_finalizadas and idx.get("estado") is not None and normalize(estado) not in pagados:
            continue
        out.append(Rendicion(str(rid), nombre, ceco, concepto, parse_date(get("fecha")),
                             parse_amount(get("monto")), _cell_link(row[idx["adjunto"]]),
                             "SI" if normalize(nombre) in cabify else "NO", estado))
        if limite and len(out) >= limite:
            break
    wb.close()
    log.info("Rendiciones a procesar: %d (con link: %d)", len(out), sum(1 for r in out if r.adjunto))
    return out


# =====================================================================
# 4. DESCARGA + DETECCIÓN DE TIPO
# =====================================================================
_session = requests.Session()
_session.headers["User-Agent"] = "Mozilla/5.0 (RPA-Rendiciones)"


def filename_from_url(url: str) -> str:
    name = unquote(Path(urlparse(url).path).name) if url.lower().startswith("http") else Path(url).name
    return name or "adjunto"


def url_key(url: str) -> str:
    return hashlib.sha1(url.encode("utf-8")).hexdigest()[:16]


def download(url: str, cache_dir: Path) -> Path:
    cache_dir.mkdir(parents=True, exist_ok=True)
    dest = cache_dir / f"{url_key(url)}__{filename_from_url(url)[:80]}"
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    if not url.lower().startswith("http"):
        shutil.copy(url, dest)
        return dest
    cfg = CONFIG["descarga"]
    last = None
    for attempt in range(1, cfg["reintentos"] + 1):
        try:
            r = _session.get(url, timeout=cfg["timeout_seg"])
            if r.status_code in (403, 404):
                raise FileNotFoundError(f"HTTP {r.status_code}: el archivo no existe o expiró el link")
            r.raise_for_status()
            dest.write_bytes(r.content)
            return dest
        except FileNotFoundError:
            raise
        except Exception as e:  # noqa: BLE001
            last = e
            time.sleep(2 * attempt)
    raise ConnectionError(f"No se pudo descargar ({cfg['reintentos']} intentos): {last}")


def sniff(data: bytes, name: str = "") -> str:
    """Tipo real por firma del archivo (la extensión puede mentir)."""
    h = data[:16]
    ext = Path(name).suffix.lower().lstrip(".")
    if h.startswith(b"%PDF"):
        return "pdf"
    if h.startswith(b"Rar!\x1a\x07"):
        return "rar"
    if h.startswith(b"7z\xbc\xaf\x27\x1c"):
        return "7z"
    if h.startswith(b"\x1f\x8b"):
        return "gz"
    if h.startswith((b"PK\x03\x04", b"PK\x05\x06")):
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            return "zip"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        if any(n.startswith("word/") for n in names):
            return "docx"
        return "zip"
    if h.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls" if ext in ("xls", "") else ext
    if h.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if h.startswith(b"\x89PNG"):
        return "png"
    if h.startswith((b"GIF87a", b"GIF89a")):
        return "gif"
    if h[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    if data[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1", b"ftypmsf1", b"ftypheif"):
        return "heic"
    if h.startswith((b"II*\x00", b"MM\x00*")):
        return "tiff"
    if h.startswith(b"BM"):
        return "bmp"
    if ext in ("csv", "txt", "tsv"):
        return ext
    if b"<html" in data[:300].lower():
        return "html"
    return ext or "desconocido"


# =====================================================================
# 5. DESCOMPRESIÓN RECURSIVA (zip, rar, 7z, gz)
# =====================================================================
@dataclass
class InnerFile:
    ruta: str
    nombre: str
    data: bytes
    tipo: str


class ArchiveError(Exception):
    pass


def _zip_members(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            if info.flag_bits & 0x1:
                raise ArchiveError("ZIP protegido con contraseña")
            name = info.filename
            if not info.flag_bits & 0x800:                      # tildes en zips de Windows
                try:
                    name = name.encode("cp437").decode("utf-8")
                except (UnicodeEncodeError, UnicodeDecodeError):
                    name = name.encode("cp437", "replace").decode("cp850", "replace")
            yield name, z.read(info)


def _rar_members(data: bytes):
    tool = find_7zip()
    if tool:                                    # 7-Zip (o UnRAR) abre RAR4 y RAR5
        yield from _extract_with_7zip(tool, data, ".rar")
        return
    import rarfile
    with tempfile.NamedTemporaryFile(suffix=".rar", delete=False) as tmp:
        tmp.write(data)
        path = tmp.name
    try:
        with rarfile.RarFile(path) as rf:
            if rf.needs_password():
                raise ArchiveError("RAR protegido con contraseña")
            for info in rf.infolist():
                if not info.is_dir():
                    yield info.filename, rf.read(info)
    except rarfile.RarCannotExec as e:
        raise ArchiveError("No se puede abrir .rar: instala 7-Zip (https://www.7-zip.org)") from e
    except rarfile.Error as e:
        raise ArchiveError(f"RAR dañado: {e}") from e
    finally:
        Path(path).unlink(missing_ok=True)


def _extract_with_7zip(tool: str, data: bytes, suffix: str):
    with tempfile.TemporaryDirectory() as tmpdir:
        src = Path(tmpdir) / f"adjunto{suffix}"
        src.write_bytes(data)
        out_dir = Path(tmpdir) / "x"
        # -pNOPASS evita que 7-Zip se quede esperando contraseña
        cp = subprocess.run([tool, "x", "-y", "-pNOPASS", f"-o{out_dir}", str(src)], capture_output=True,
                            text=True, errors="replace", creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        if cp.returncode != 0:
            msg = (cp.stderr or cp.stdout).lower()
            if "password" in msg or "contrase" in msg:
                raise ArchiveError("Comprimido protegido con contraseña")
            raise ArchiveError(f"7-Zip no pudo abrir el comprimido: {(cp.stderr or cp.stdout).strip()[-160:]}")
        for p in out_dir.rglob("*"):
            if p.is_file():
                yield str(p.relative_to(out_dir)), p.read_bytes()


def _7z_members(data: bytes):
    import py7zr
    with tempfile.TemporaryDirectory() as tmpdir:
        with py7zr.SevenZipFile(io.BytesIO(data), mode="r") as z:
            if z.needs_password():
                raise ArchiveError("7z protegido con contraseña")
            z.extractall(path=tmpdir)
        for p in Path(tmpdir).rglob("*"):
            if p.is_file():
                yield str(p.relative_to(tmpdir)), p.read_bytes()


_SKIP = ("__macosx/", "thumbs.db", ".ds_store", "desktop.ini")


def expand(name: str, data: bytes, depth: int = 0, prefix: str = "") -> list[InnerFile]:
    lim = CONFIG["descompresion"]
    tipo = sniff(data, name)
    ruta = f"{prefix}{name}"
    if tipo not in ("zip", "rar", "7z", "gz"):
        return [InnerFile(ruta, Path(name).name, data, tipo)]
    if depth >= lim["profundidad_max"]:
        raise ArchiveError(f"Demasiados niveles de compresión en {ruta}")
    try:
        members = {"zip": _zip_members, "rar": _rar_members, "7z": _7z_members}.get(tipo)
        members = members(data) if members else [(Path(name).stem or "contenido", gzip.decompress(data))]
        out: list[InnerFile] = []
        for mname, mdata in members:
            low = mname.lower().replace("\\", "/")
            if any(s in low for s in _SKIP) or Path(mname).name.startswith("._") or not mdata:
                continue
            out.extend(expand(mname, mdata, depth + 1, f"{ruta} > "))
            if len(out) > lim["max_archivos_por_adjunto"]:
                log.warning("%s: demasiados archivos, se truncó", ruta)
                break
        return out
    except zipfile.BadZipFile as e:
        raise ArchiveError(f"Comprimido dañado: {ruta}") from e


# =====================================================================
# 6. EXTRACCIÓN: archivo -> unidades (texto y/o imagen de página)
# =====================================================================
IMAGE_TYPES = {"jpg", "jpeg", "png", "gif", "webp", "heic", "tiff", "bmp"}
SHEET_TYPES = {"xlsx", "xlsm", "xls", "csv", "tsv"}


@dataclass
class Unidad:
    archivo: InnerFile
    pagina: int
    origen: str
    texto: str = ""
    imagen: bytes | None = None
    filas: list = field(default_factory=list)


class UnsupportedFile(Exception):
    pass


def _to_jpeg(img: Image.Image, max_side: int = 2200) -> bytes:
    img = ImageOps.exif_transpose(img)
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    if max(img.size) > max_side:
        img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=88)
    return buf.getvalue()


def extract(f: InnerFile) -> list[Unidad]:
    dcfg = CONFIG["documentos"]
    if f.tipo == "pdf":
        import pymupdf
        try:
            doc = pymupdf.open(stream=f.data, filetype="pdf")
        except Exception as e:  # noqa: BLE001
            raise UnsupportedFile(f"PDF dañado o ilegible: {e}") from e
        if doc.needs_pass:
            raise UnsupportedFile("PDF protegido con contraseña")
        out = []
        for i, page in enumerate(doc):
            if i >= dcfg["max_paginas_pdf"]:
                break
            text = page.get_text("text") or ""
            if len(text.strip()) >= dcfg["min_chars_pdf_con_texto"]:
                out.append(Unidad(f, i + 1, "pdf_texto", texto=text))
            else:                                        # escaneado -> imagen
                pix = page.get_pixmap(dpi=CONFIG["local"]["dpi_pdf_escaneado"])
                out.append(Unidad(f, i + 1, "pdf_escaneado", imagen=_to_jpeg(Image.open(io.BytesIO(pix.tobytes("png"))))))
        doc.close()
        return out

    if f.tipo in IMAGE_TYPES:
        try:
            img = Image.open(io.BytesIO(f.data))
            frames = []
            for i in range(getattr(img, "n_frames", 1)):
                img.seek(i)
                frames.append(img.copy())
        except Exception as e:  # noqa: BLE001
            raise UnsupportedFile(f"Imagen ilegible ({f.tipo}): {e}") from e
        return [Unidad(f, i + 1, "imagen", imagen=_to_jpeg(fr)) for i, fr in enumerate(frames)]

    if f.tipo in SHEET_TYPES:
        import pandas as pd
        bio = io.BytesIO(f.data)
        try:
            if f.tipo in ("csv", "tsv"):
                sheets = {"csv": pd.read_csv(bio, sep="\t" if f.tipo == "tsv" else None, engine="python",
                                             header=None, dtype=object, encoding_errors="replace")}
            else:
                sheets = pd.read_excel(bio, sheet_name=None, header=None, dtype=object,
                                       engine="xlrd" if f.tipo == "xls" else "openpyxl")
        except Exception as e:  # noqa: BLE001
            raise UnsupportedFile(f"Planilla ilegible: {e}") from e
        out = []
        if len(sheets) > 1:                             # evita doble conteo detalle + resumen
            detalle = {k: v for k, v in sheets.items()
                       if not re.search(r"resumen|total|consolidado|dashboard|grafico", normalize(k))}
            sheets = detalle or sheets
        for i, (sname, df) in enumerate(sheets.items()):
            df = df.dropna(how="all").dropna(axis=1, how="all")
            if df.empty:
                continue
            filas = [[None if pd.isna(v) else v for v in row] for row in df.itertuples(index=False)]
            texto = "\n".join("\t".join("" if v is None else str(v) for v in r) for r in filas)
            out.append(Unidad(f, i + 1, "planilla", texto=f"[Hoja: {sname}]\n{texto}", filas=filas))
        return out

    if f.tipo == "docx":
        import docx
        try:
            d = docx.Document(io.BytesIO(f.data))
        except Exception as e:  # noqa: BLE001
            raise UnsupportedFile(f"Word ilegible: {e}") from e
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        parts += ["\t".join(c.text.strip() for c in row.cells) for t in d.tables for row in t.rows]
        out = [Unidad(f, 1, "word", texto="\n".join(parts))] if parts else []
        for j, rel in enumerate(r for r in d.part.rels.values() if "image" in r.reltype):
            try:
                out.append(Unidad(f, j + 2, "imagen", imagen=_to_jpeg(Image.open(io.BytesIO(rel.target_part.blob)))))
            except Exception:  # noqa: BLE001
                continue
        return out

    if f.tipo == "txt":
        return [Unidad(f, 1, "texto", texto=f.data.decode("utf-8", "replace"))]
    if f.tipo == "html":
        raise UnsupportedFile("El link devolvió una página web (archivo borrado o requiere login)")
    raise UnsupportedFile(f"Tipo de archivo no soportado: {f.tipo}")


# =====================================================================
# 7. CLASIFICACIÓN Y ANÁLISIS
# =====================================================================
class Classifier:
    def __init__(self, conceptos: dict):
        self.conceptos = list(conceptos)
        self._rules = [(c, normalize(kw), re.compile(rf"(?<![a-z0-9]){re.escape(normalize(kw))}(?![a-z0-9])"))
                       for c, kws in conceptos.items() for kw in kws]

    def classify(self, *texts: str) -> tuple[str, str]:
        t = normalize(" ".join(x for x in texts if x))
        for c, kw, rx in self._rules:
            if rx.search(t):
                return c, kw
        return "OTROS", ""

    def valid(self, c) -> str:
        c = (c or "").strip().upper()
        return c if c in self.conceptos else "OTROS"


@dataclass
class Item:
    tipo_documento: str = ""
    emisor: str = ""
    rut_emisor: str | None = None
    folio: str | None = None
    fecha_doc: object = None
    descripcion: str = ""
    monto: float | None = None
    concepto: str = "OTROS"
    palabra_clave: str = ""
    metodo: str = ""
    confianza: str = ""
    observacion: str = ""


_TOTAL_RX = re.compile(r"(total\s+a\s+pagar|monto\s+total|total\s+pagado|valor\s+total|\btotal\b)")


class LocalAnalyzer:
    """Texto nativo + OCR Tesseract + reglas. Gratis; nada sale del equipo."""
    name = "local"

    def __init__(self, clf: Classifier):
        self.clf = clf
        self.lang = CONFIG["local"]["idioma_ocr"]
        self._tess = None
        try:
            import pytesseract
            cmd = find_tesseract()
            if cmd:
                pytesseract.pytesseract.tesseract_cmd = cmd
            langs = pytesseract.get_languages(config="")
            if self.lang not in langs:
                log.warning("Tesseract no tiene idioma español instalado; se usará inglés (menos preciso).")
                self.lang = "eng"
            self._tess = pytesseract
        except Exception:  # noqa: BLE001
            log.warning("Tesseract NO está instalado: fotos y PDF escaneados no se podrán leer en motor Local. "
                        "Instálalo o usa el motor Claude.")

    def _from_text(self, text: str, metodo: str) -> list[Item]:
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        total, conf = None, "baja"
        for i, ln in enumerate(lines):
            nl = normalize(ln)
            if _TOTAL_RX.search(nl) and "subtotal" not in nl and "sub total" not in nl:
                vals = find_amounts(ln) or (find_amounts(lines[i + 1]) if i + 1 < len(lines) else [])
                if vals:
                    total, conf = max(vals) if total is None else max(total, max(vals)), "media"
        if total is None:
            vals = find_amounts(_RUT_RE.sub(" ", text), max_v=5_000_000)
            total = max(vals) if vals else None
        emisor = next((ln for ln in lines[:6] if re.search(r"[A-Za-z]{3,}", ln)), "")[:80]
        c, kw = self.clf.classify(text)
        return [Item("boleta/factura" if re.search(r"boleta|factura", normalize(text)) else "documento",
                     emisor, find_rut(text), find_folio(text), find_date(text),
                     (f"[{kw}] " if kw else "") + " ".join(lines[:3])[:150], total, c, kw, metodo, conf,
                     "" if total else "No se detectó monto")]

    def _from_sheet(self, filas: list) -> list[Item]:
        hdr, cm, cd, cf = None, None, None, None
        for i, row in enumerate(filas[:15]):
            cells = [normalize(v) if isinstance(v, str) else "" for v in row]
            m = next((j for j, c in enumerate(cells) if c and any(h in c for h in
                      ("monto", "total", "valor", "importe", "precio", "costo", "$"))), None)
            if m is not None:
                hdr, cm = i, m
                cd = next((j for j, c in enumerate(cells) if c and any(h in c for h in
                           ("descripcion", "detalle", "concepto", "glosa", "motivo", "item", "proveedor", "gasto"))), None)
                cf = next((j for j, c in enumerate(cells) if "fecha" in c), None)
                break
        items = []
        for row in filas[(hdr + 1) if hdr is not None else 0:]:
            joined = " ".join(str(v) for v in row if isinstance(v, str) and v.strip())
            if re.search(r"\b(sub ?)?total(es)?\b|resumen", normalize(joined)):
                if hdr is not None:
                    break                                  # fin de la tabla: lo de abajo suele ser resumen/firmas
                continue
            if cm is not None and cm < len(row):
                monto = parse_amount(row[cm])
            else:
                nums = [parse_amount(v) for v in row if isinstance(v, (int, float)) and not isinstance(v, bool)]
                nums = [n for n in nums if n and 100 <= n <= 50_000_000]
                monto = nums[-1] if nums else None
            if not monto or monto <= 0:
                continue
            desc = str(row[cd]) if cd is not None and cd < len(row) and row[cd] else joined
            c, kw = self.clf.classify(joined)
            items.append(Item("planilla", "", None, None, parse_date(row[cf]) if cf is not None and cf < len(row) else None,
                              desc[:150], monto, c, kw, "planilla_reglas", "alta" if cm is not None else "media"))
        return items

    def analyze(self, u: Unidad) -> list[Item]:
        if u.origen == "planilla":
            return self._from_sheet(u.filas)
        if u.imagen is not None:
            if not self._tess:
                return [Item(metodo="sin_ocr", confianza="baja", observacion="Imagen no leída: falta Tesseract")]
            text = self._tess.image_to_string(Image.open(io.BytesIO(u.imagen)).convert("L"),
                                              lang=self.lang, config="--psm 6")
            if len(text.strip()) < 15:
                return [Item(metodo="ocr_tesseract", confianza="baja", observacion="Imagen ilegible por OCR")]
            return self._from_text(text, "ocr_tesseract")
        return self._from_text(u.texto, f"texto_{u.origen}")


_SYSTEM = ("Eres un auditor de rendiciones de gastos de una empresa chilena. Recibes una página (foto, escaneo, "
           "PDF o planilla) que respalda una rendición. Extrae TODOS los comprobantes visibles (boletas, facturas, "
           "vouchers, tickets, capturas de apps, planillas de gastos). Montos en pesos chilenos: el punto separa "
           "miles ('$12.500' = 12500). Responde SOLO con JSON válido, sin texto adicional ni bloques ```.")

_PROMPT = """Clasifica cada ítem en UNO de estos conceptos: {conceptos}.
Pistas: Copec/Shell/Aramco/Enex/litros/bencina/diésel = COMBUSTIBLE; Cabify/Uber/DiDi/inDrive = APPS_TRANSPORTE;
autopistas/TAG/peajes = PEAJE_TAG. Si una planilla lista gastos, cada fila es un ítem.
Si un comprobante sólo muestra el total, crea un único ítem con ese total. No inventes montos.
Formato:
{{"documentos": [{{"tipo_documento": "boleta|factura|voucher|ticket|captura_app|planilla|manuscrito|otro",
  "emisor": "comercio", "rut_emisor": "12345678-9 o null", "folio": "número o null", "fecha": "YYYY-MM-DD o null",
  "monto_total": 12345, "items": [{{"descripcion": "texto breve", "monto": 12345, "concepto": "CONCEPTO"}}],
  "legible": true, "observacion": "enmendaduras, manuscrito, fecha ilegible, etc. o vacío"}}]}}
Si la página no contiene comprobantes, devuelve {{"documentos": []}}."""


class ClaudeAnalyzer:
    """Claude visión: muy superior en fotos, escaneos, boletas térmicas y manuscritos."""
    name = "claude"

    def __init__(self, clf: Classifier, api_key: str):
        import anthropic
        self.client = anthropic.Anthropic(api_key=api_key, max_retries=5)
        self.cfg = CONFIG["claude"]
        self.clf = clf
        self.prompt = _PROMPT.format(conceptos=", ".join(clf.conceptos))

    def analyze(self, u: Unidad) -> list[Item]:
        content = []
        if u.imagen is not None:
            img = Image.open(io.BytesIO(u.imagen))
            img_bytes = u.imagen
            if max(img.size) > self.cfg["lado_max_imagen_px"]:
                img_bytes = _to_jpeg(img, self.cfg["lado_max_imagen_px"])
            content.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg",
                                                        "data": base64.b64encode(img_bytes).decode()}})
        if u.texto:
            content.append({"type": "text", "text": f"Archivo '{u.archivo.nombre}':\n{u.texto[:self.cfg['max_chars_texto']]}"})
        content.append({"type": "text", "text": self.prompt})
        resp = self.client.messages.create(model=self.cfg["modelo"], max_tokens=self.cfg["max_tokens"],
                                           system=_SYSTEM, messages=[{"role": "user", "content": content}])
        raw = "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
        try:
            clean = re.sub(r"```(?:json)?", "", raw)
            data = json.loads(clean[clean.find("{"):clean.rfind("}") + 1])
        except (ValueError, json.JSONDecodeError):
            return [Item(metodo=f"claude_{u.origen}", confianza="baja", observacion=f"Respuesta no interpretable: {raw[:100]}")]
        items = []
        for d in data.get("documentos", []):
            base = dict(tipo_documento=d.get("tipo_documento") or "", emisor=(d.get("emisor") or "")[:80],
                        rut_emisor=d.get("rut_emisor"), folio=str(d["folio"]) if d.get("folio") else None,
                        fecha_doc=parse_date(d.get("fecha")), metodo=f"claude_{u.origen}",
                        confianza="alta" if d.get("legible", True) else "baja", observacion=d.get("observacion") or "")
            for it in d.get("items") or [{"descripcion": base["emisor"], "monto": d.get("monto_total")}]:
                c = self.clf.valid(it.get("concepto"))
                rc, kw = self.clf.classify(it.get("descripcion", ""), base["emisor"])
                items.append(Item(**base, descripcion=(it.get("descripcion") or "")[:150],
                                  monto=parse_amount(it.get("monto")), concepto=rc if c == "OTROS" else c,
                                  palabra_clave=kw))
        return items or [Item(metodo=f"claude_{u.origen}", confianza="media", observacion="Página sin comprobantes")]


# =====================================================================
# 8. PROCESO POR RENDICIÓN (con caché)
# =====================================================================
def _json_default(o):
    if isinstance(o, (dt.date, dt.datetime)):
        return o.isoformat()
    raise TypeError(type(o))


_SIN_OCR = "falta Tesseract"


def _incompleto(res: dict) -> bool:
    """Resultados que dependen de algo no instalado: no se guardan en caché para reintentarlos."""
    return res.get("estado") == "ERROR COMPRIMIDO" or any(
        _SIN_OCR in (it.get("observacion") or "") for it in res.get("items", []))


def process_one(r: Rendicion, analyzer, cache: Path, reprocess: bool) -> dict:
    res = {"rendicion": asdict(r), "estado": "OK", "archivos": [], "items": [], "errores": []}
    if not r.adjunto:
        res["estado"] = "SIN ADJUNTO"
        return res
    res_file = cache / "resultados" / f"{url_key(r.adjunto)}_{analyzer.name}.json"
    if res_file.exists() and not reprocess:
        cached = json.loads(res_file.read_text(encoding="utf-8"))
        if not _incompleto(cached):                     # resultados de corridas sin Tesseract/7-Zip se rehacen
            cached["rendicion"] = asdict(r)
            return cached

    try:
        data = download(r.adjunto, cache / "descargas").read_bytes()
        res["hash_adjunto"] = sha256(data)
        inner = expand(filename_from_url(r.adjunto), data)
    except (FileNotFoundError, ConnectionError) as e:
        res["estado"], res["errores"] = "ERROR DESCARGA", [str(e)]
        return res                                          # no se guarda en caché: se reintenta
    except ArchiveError as e:
        res["estado"], res["errores"], inner = "ERROR COMPRIMIDO", [str(e)], []

    for f in inner:
        finfo = {"ruta": f.ruta, "tipo": f.tipo, "hash": sha256(f.data)}
        try:
            for u in extract(f):
                try:
                    for it in analyzer.analyze(u):
                        res["items"].append({**asdict(it), "archivo": f.ruta, "tipo_archivo": f.tipo, "pagina": u.pagina})
                except Exception as e:  # noqa: BLE001
                    res["errores"].append(f"{f.ruta} p.{u.pagina}: {e}")
                    res["items"].append({**asdict(Item(metodo=analyzer.name, confianza="baja",
                                                       observacion=f"Error al analizar: {str(e)[:150]}")),
                                         "archivo": f.ruta, "tipo_archivo": f.tipo, "pagina": u.pagina})
        except UnsupportedFile as e:
            res["errores"].append(f"{f.ruta}: {e}")
        res["archivos"].append(finfo)

    sin_ocr = sum(1 for it in res["items"] if _SIN_OCR in (it.get("observacion") or ""))
    if sin_ocr:
        res["errores"].append(f"{sin_ocr} imagen(es)/escaneado(s) sin leer: falta instalar Tesseract")
    if res["estado"] == "OK":
        if not res["items"]:
            res["estado"] = "SIN CONTENIDO LEGIBLE"
        elif sin_ocr and not any(it.get("monto") for it in res["items"]):
            res["estado"] = "SIN LEER (FALTA OCR)"
        elif res["errores"]:
            res["estado"] = "PARCIAL"
    if not _incompleto(res):
        res_file.parent.mkdir(parents=True, exist_ok=True)
        res_file.write_text(json.dumps(res, ensure_ascii=False, default=_json_default), encoding="utf-8")
    return res


def _safe(name: str, maxlen: int = 60) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name)).strip(" .")
    return (name or "sin_nombre")[:maxlen]


def download_only(rendiciones: list[Rendicion], dest_root: Path, cache: Path,
                  progress=None, stop_event: threading.Event | None = None) -> Path:
    """Descarga, descomprime y ordena los adjuntos en carpetas CECO/Rendidor/, con un índice CSV.
    Pensado para revisarlos luego con Claude Cowork (o cualquier herramienta) sin depender de Tesseract/API."""
    import csv
    dest_root.mkdir(parents=True, exist_ok=True)
    index_rows, total = [], len(rendiciones)

    def one(r: Rendicion):
        rows = []
        if not r.adjunto:
            return rows, "SIN ADJUNTO"
        try:
            data = download(r.adjunto, cache / "descargas").read_bytes()
            inner = expand(filename_from_url(r.adjunto), data)
        except (FileNotFoundError, ConnectionError) as e:
            return rows, f"ERROR DESCARGA: {e}"
        except ArchiveError as e:
            return rows, f"ERROR COMPRIMIDO: {e}"
        folder = dest_root / _safe(r.ceco) / _safe(r.rendidor)
        folder.mkdir(parents=True, exist_ok=True)
        for k, f in enumerate(inner, 1):
            ruta_interna = f.ruta.split(" > ", 1)[1] if " > " in f.ruta else f.nombre
            fname = f"{r.id}__{k:02d}__{_safe(Path(ruta_interna).stem, 50)}{Path(f.nombre).suffix or '.' + f.tipo}"
            (folder / fname).write_bytes(f.data)
            rows.append([r.id, r.ceco, r.rendidor, r.concepto_declarado, r.fecha, r.monto_declarado, r.estado,
                         str((folder / fname).relative_to(dest_root)), f.tipo, f.ruta])
        return rows, "OK"

    done = 0
    with ThreadPoolExecutor(max_workers=CONFIG["descarga"]["workers"]) as ex:
        futs = {ex.submit(one, r): r for r in rendiciones}
        for fut in as_completed(futs):
            r = futs[fut]
            try:
                rows, st = fut.result()
            except Exception as e:  # noqa: BLE001
                rows, st = [], f"ERROR: {e}"
            if not rows:
                index_rows.append([r.id, r.ceco, r.rendidor, r.concepto_declarado, r.fecha, r.monto_declarado,
                                   r.estado, "", "", st])
            index_rows.extend(rows)
            done += 1
            log.info("[%d/%d] %s | %s | %s | %d archivo(s)", done, total, r.id, r.rendidor[:25], st, len(rows))
            if progress:
                progress(done, total)
            if stop_event is not None and stop_event.is_set():
                log.warning("Detenido por el usuario.")
                for f in futs:
                    f.cancel()
                break

    idx_path = dest_root / "indice_adjuntos.csv"
    with open(idx_path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["ID", "CECO", "Rendidor", "Cuenta contable", "Fecha rendicion", "Monto declarado", "Estado",
                    "Archivo (relativo)", "Tipo", "Origen / estado"])
        w.writerows(index_rows)
    log.info("Adjuntos ordenados en: %s", dest_root)
    log.info("Índice: %s", idx_path)
    return dest_root


# =====================================================================
# 9. AUDITORÍA
# =====================================================================
def audit(results: list[dict]) -> None:
    acfg = CONFIG["auditoria"]
    alert_c, tol, max_days = set(acfg["conceptos_alerta"]), acfg["tolerancia_diferencia_pct"] / 100, acfg["dias_max_antiguedad_doc"]
    by_adj, by_file, by_folio, by_trip = (defaultdict(set) for _ in range(4))
    trip = lambda it: (normalize(it["emisor"])[:25], str(it["fecha_doc"]), it["monto"])  # noqa: E731
    for res in results:
        rid = res["rendicion"]["id"]
        if res.get("hash_adjunto"):
            by_adj[res["hash_adjunto"]].add(rid)
        for f in res["archivos"]:
            by_file[f["hash"]].add(rid)
        for it in res["items"]:
            if it.get("rut_emisor") and it.get("folio"):
                by_folio[(it["rut_emisor"], str(it["folio"]))].add(rid)
            elif it.get("monto") and it.get("fecha_doc") and it.get("emisor"):
                by_trip[trip(it)].add(rid)

    for res in results:
        r, rid = res["rendicion"], res["rendicion"]["id"]
        f_rend, decl = parse_date(r.get("fecha")), normalize(r.get("concepto_declarado"))
        for it in res["items"]:
            a, c = [], it.get("concepto") or "OTROS"
            if c in alert_c:
                a.append(f"ALERTA {c}")
                if c == "APPS_TRANSPORTE" and r.get("usa_cabify") == "SI":
                    a.append("Rinde app de transporte y además es usuario Cabify corporativo")
            if c == "COMBUSTIBLE" and "movilizacion" in decl:
                a.append("Combustible rendido como Movilización")
            fd = parse_date(it.get("fecha_doc"))
            if fd and f_rend:
                days = (f_rend - fd).days
                if days > max_days:
                    a.append(f"Documento antiguo ({days} días antes de la rendición)")
                elif days < -1:
                    a.append("Fecha del documento posterior a la rendición")
            if it.get("rut_emisor") and it.get("folio"):
                ids = by_folio[(it["rut_emisor"], str(it["folio"]))] - {rid}
                if ids:
                    a.append(f"Comprobante (RUT+folio) repetido en ID {', '.join(sorted(ids))}")
            elif it.get("monto") and it.get("fecha_doc") and it.get("emisor"):
                ids = by_trip[trip(it)] - {rid}
                if ids:
                    a.append(f"Posible duplicado (emisor+fecha+monto) en ID {', '.join(sorted(ids))}")
            it["alertas"] = "; ".join(a)

        r_alerts = []
        if res.get("hash_adjunto") and (ids := by_adj[res["hash_adjunto"]] - {rid}):
            r_alerts.append(f"Mismo adjunto usado en ID {', '.join(sorted(ids))}")
        dup = set().union(*[by_file[f["hash"]] - {rid} for f in res["archivos"]]) if res["archivos"] else set()
        if dup and not r_alerts:
            r_alerts.append(f"Archivo interno repetido en ID {', '.join(sorted(dup))}")
        detected = sum(it["monto"] for it in res["items"] if it.get("monto"))
        res["monto_detectado"] = detected or None
        declared = r.get("monto_declarado")
        if declared and detected and detected > 2 * declared:
            r_alerts.append("Monto detectado > 2x lo declarado: revisar (posible doble conteo en planilla)")
        elif declared and detected and abs(detected - declared) > tol * declared:
            r_alerts.append(f"Diferencia de monto {detected - declared:+,.0f}".replace(",", "."))
        elif declared and not detected and res["estado"] in ("OK", "PARCIAL"):
            r_alerts.append("No se pudo detectar monto en los respaldos")
        res["conceptos_detectados"] = ", ".join(sorted({it["concepto"] for it in res["items"] if it.get("monto")}))
        item_alerts = sorted({x for it in res["items"] for x in (it.get("alertas") or "").split("; ") if x})
        res["alertas"] = "; ".join(r_alerts + item_alerts)


SIN_ID = "SIN_IDENTIFICAR"


def add_cuadratura(results: list[dict]) -> None:
    """Agrega un ítem SIN_IDENTIFICAR por la diferencia declarado - detectado (si es positiva),
    para que los totales por concepto cuadren con lo rendido en Talana."""
    for res in results:
        declared = res["rendicion"].get("monto_declarado") or 0
        detected = sum(it["monto"] for it in res["items"] if it.get("monto") and it.get("concepto") != SIN_ID)
        gap = declared - detected
        if gap > 1:
            motivo = {"SIN ADJUNTO": "Rendición sin adjunto",
                      "ERROR DESCARGA": "Adjunto no descargable",
                      "ERROR COMPRIMIDO": "Comprimido no se pudo abrir",
                      "SIN LEER (FALTA OCR)": "Imagen/escaneado sin leer (falta OCR)"}.get(
                res["estado"], "Parte del monto declarado no se identificó en los respaldos")
            res["items"].append({**asdict(Item(descripcion=motivo, monto=gap, concepto=SIN_ID,
                                               metodo="cuadratura", confianza="")),
                                 "archivo": "", "tipo_archivo": "", "pagina": None, "alertas": ""})


# =====================================================================
# 10. REPORTE EXCEL
# =====================================================================
MONEY = '_ * #,##0_ ;_ * \\-#,##0_ ;_ * "-"_ ;_ @_ '
HDR_FILL, HDR_FONT = PatternFill("solid", start_color="1F4E78"), Font(bold=True, color="FFFFFF")
ALERT_FILL, WARN_FILL = PatternFill("solid", start_color="FFC7CE"), PatternFill("solid", start_color="FFF2CC")
TOTAL_FILL = PatternFill("solid", start_color="D9E1F2")
BORDER = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _header(ws, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _total_row(ws, row, c_from, c_to, label_col=1):
    ws.cell(row=row, column=label_col, value="TOTAL")
    for j in range(c_from, c_to + 1):
        col = get_column_letter(j)
        ws.cell(row=row, column=j, value=f"=SUM({col}2:{col}{row - 1})").number_format = MONEY
    for j in range(1, c_to + 1):
        ws.cell(row=row, column=j).fill, ws.cell(row=row, column=j).font = TOTAL_FILL, Font(bold=True)


def write_report(results: list[dict], out_path: Path, meta: dict) -> Path:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Resumen x Rendición"
    h = ["ID", "Rendidor", "Usuario Cabify", "CECO", "Concepto declarado", "Fecha rendición", "Monto declarado",
         "Monto detectado", "Diferencia", "N° archivos", "N° ítems", "Conceptos detectados", "Estado", "Alertas", "Adjunto"]
    _header(ws, h)
    for i, res in enumerate(results, 2):
        r = res["rendicion"]
        vals = [r["id"], r["rendidor"], r["usa_cabify"], r["ceco"], r["concepto_declarado"], parse_date(r.get("fecha")),
                r.get("monto_declarado"), res.get("monto_detectado"),
                f'=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i})),H{i}-G{i},"")', len(res["archivos"]), len(res["items"]),
                res.get("conceptos_detectados", ""), res["estado"], res.get("alertas", ""), "Abrir" if r.get("adjunto") else ""]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v)
        ws.cell(row=i, column=6).number_format = "DD-MM-YYYY"
        for c in (7, 8, 9):
            ws.cell(row=i, column=c).number_format = MONEY
        if r.get("adjunto"):
            ws.cell(row=i, column=15).hyperlink = r["adjunto"]
            ws.cell(row=i, column=15).font = Font(color="0563C1", underline="single")
        fill = ALERT_FILL if res.get("alertas") else (WARN_FILL if res["estado"] != "OK" else None)
        if fill:
            for c in range(1, len(h) + 1):
                ws.cell(row=i, column=c).fill = fill
    _widths(ws, [11, 28, 9, 26, 24, 12, 13, 13, 13, 8, 8, 30, 18, 60, 9])
    ws.auto_filter.ref = f"A1:O{max(2, len(results) + 1)}"

    wd = wb.create_sheet("Desglose x Ítem")
    hd = ["ID", "Rendidor", "Usuario Cabify", "CECO", "Concepto declarado", "Mes rendición", "Archivo (ruta interna)",
          "Tipo archivo", "Página", "Método lectura", "Tipo documento", "Emisor", "RUT emisor", "Folio",
          "Fecha documento", "Descripción", "Monto ítem", "Concepto detectado", "Palabra clave", "Confianza",
          "Alertas", "Observación"]
    _header(wd, hd)
    n = 1
    for res in results:
        r = res["rendicion"]
        fr = parse_date(r.get("fecha"))
        for it in res["items"]:
            n += 1
            vals = [r["id"], r["rendidor"], r["usa_cabify"], r["ceco"], r["concepto_declarado"],
                    fr.strftime("%Y-%m") if fr else "", it.get("archivo"), it.get("tipo_archivo"), it.get("pagina"),
                    it.get("metodo"), it.get("tipo_documento"), it.get("emisor"), it.get("rut_emisor"), it.get("folio"),
                    parse_date(it.get("fecha_doc")), it.get("descripcion"), it.get("monto"), it.get("concepto"),
                    it.get("palabra_clave"), it.get("confianza"), it.get("alertas"), it.get("observacion")]
            for c, v in enumerate(vals, 1):
                wd.cell(row=n, column=c, value=v)
            wd.cell(row=n, column=15).number_format = "DD-MM-YYYY"
            wd.cell(row=n, column=17).number_format = MONEY
            if it.get("alertas"):
                for c in range(1, len(hd) + 1):
                    wd.cell(row=n, column=c).fill = ALERT_FILL
    last = max(n, 2)
    _widths(wd, [11, 26, 9, 24, 22, 10, 40, 8, 7, 18, 14, 26, 12, 10, 12, 40, 12, 20, 14, 9, 50, 35])
    wd.auto_filter.ref = f"A1:V{last}"

    D = "'Desglose x Ítem'"
    R_REND, R_CONC, R_MONTO, R_MES = (f"{D}!${c}$2:${c}${last}" for c in ("B", "R", "Q", "F"))
    used = [c for c in CONFIG["conceptos"] if any(it.get("concepto") == c for res in results for it in res["items"])] \
        or list(CONFIG["conceptos"])
    if any(it.get("concepto") == SIN_ID for res in results for it in res["items"]):
        used.append(SIN_ID)
    GREY_FILL = PatternFill("solid", start_color="E7E6E6")

    wp = wb.create_sheet("Rendidor x Concepto")
    _header(wp, ["Rendidor", "Usuario Cabify"] + used + ["Total detectado"])
    people = sorted({(res["rendicion"]["rendidor"], res["rendicion"]["usa_cabify"]) for res in results})
    tc = len(used) + 3
    for i, (name, cab) in enumerate(people, 2):
        wp.cell(row=i, column=1, value=name)
        wp.cell(row=i, column=2, value=cab)
        for j in range(3, tc):
            col = get_column_letter(j)
            wp.cell(row=i, column=j, value=f"=SUMIFS({R_MONTO},{R_REND},$A{i},{R_CONC},{col}$1)").number_format = MONEY
        wp.cell(row=i, column=tc, value=f"=SUM(C{i}:{get_column_letter(tc - 1)}{i})").number_format = MONEY
    _total_row(wp, len(people) + 2, 3, tc)
    for j, c in enumerate(used, 3):
        if c in CONFIG["auditoria"]["conceptos_alerta"]:
            for rr in range(2, len(people) + 2):
                wp.cell(row=rr, column=j).font = Font(bold=True, color="C00000")
    _widths(wp, [28, 9] + [14] * (len(used) + 1))

    # CECO x Rendidor x Concepto detectado (mes a mes, fórmulas sobre Desglose)
    wc = wb.create_sheet("CECO x Rendidor x Concepto")
    R_CECO = f"{D}!$D$2:$D${last}"
    meses_c = sorted({parse_date(res["rendicion"].get("fecha")).strftime("%Y-%m")
                      for res in results if parse_date(res["rendicion"].get("fecha"))})
    hc = ["CECO", "Rendidor", "Concepto detectado (adjuntos)"] + meses_c + ["Total", "% del rendidor"]
    _header(wc, hc)
    wc.freeze_panes = "D2"
    combos_c = defaultdict(float)
    for res in results:
        rr_ = res["rendicion"]
        for it in res["items"]:
            if it.get("monto"):
                combos_c[(rr_["ceco"], rr_["rendidor"], it.get("concepto") or "OTROS")] += it["monto"]
    tcc = len(meses_c) + 4
    row_c, ceco_tot_rows = 2, []
    for ceco in sorted({k[0] for k in combos_c}):
        ceco_start, rend_tot_rows = row_c, []
        for rend in sorted({k[1] for k in combos_c if k[0] == ceco}):
            start = row_c
            for conc in sorted({k[2] for k in combos_c if k[0] == ceco and k[1] == rend},
                               key=lambda c: -combos_c[(ceco, rend, c)]):
                wc.cell(row=row_c, column=1, value=ceco)
                wc.cell(row=row_c, column=2, value=rend)
                wc.cell(row=row_c, column=3, value=conc)
                for j, m in enumerate(meses_c, 4):
                    wc.cell(row=row_c, column=j, value=(
                        f"=SUMIFS({R_MONTO},{R_CECO},$A{row_c},{R_REND},$B{row_c},{R_CONC},$C{row_c},{R_MES},\"{m}\")"
                    )).number_format = MONEY
                if conc in CONFIG["auditoria"]["conceptos_alerta"] or conc == SIN_ID:
                    for c in range(1, tcc + 2):
                        wc.cell(row=row_c, column=c).fill = ALERT_FILL if conc != SIN_ID else GREY_FILL
                row_c += 1
            end = row_c - 1
            wc.cell(row=row_c, column=1, value=ceco)
            wc.cell(row=row_c, column=2, value=rend)
            wc.cell(row=row_c, column=3, value="Total rendidor")
            for j in range(4, tcc):
                wc.cell(row=row_c, column=j, value=f"=SUM({get_column_letter(j)}{start}:{get_column_letter(j)}{end})").number_format = MONEY
            for x in range(start, end + 1):
                wc.cell(row=x, column=tcc + 1,
                        value=f"=IF($"+get_column_letter(tcc)+f"${row_c}=0,\"\",{get_column_letter(tcc)}{x}/$"
                              + get_column_letter(tcc) + f"${row_c})").number_format = "0.0%"
            for c in range(1, tcc + 2):
                wc.cell(row=row_c, column=c).fill, wc.cell(row=row_c, column=c).font = WARN_FILL, Font(bold=True)
            rend_tot_rows.append(row_c)
            row_c += 1
        wc.cell(row=row_c, column=1, value=ceco)
        wc.cell(row=row_c, column=2, value=f"TOTAL {ceco}")
        for j in range(4, tcc):
            col = get_column_letter(j)
            wc.cell(row=row_c, column=j, value="=" + "+".join(f"{col}{x}" for x in rend_tot_rows)).number_format = MONEY
        for c in range(1, tcc + 2):
            wc.cell(row=row_c, column=c).fill, wc.cell(row=row_c, column=c).font = TOTAL_FILL, Font(bold=True)
        ceco_tot_rows.append(row_c)
        row_c += 1
    if ceco_tot_rows:
        wc.cell(row=row_c, column=1, value="TOTAL GENERAL")
        for j in range(4, tcc):
            col = get_column_letter(j)
            wc.cell(row=row_c, column=j, value="=" + "+".join(f"{col}{x}" for x in ceco_tot_rows)).number_format = MONEY
        for c in range(1, tcc + 2):
            wc.cell(row=row_c, column=c).fill, wc.cell(row=row_c, column=c).font = TOTAL_FILL, Font(bold=True)
    for x in range(2, row_c + 1):
        wc.cell(row=x, column=tcc, value=f"=SUM(D{x}:{get_column_letter(tcc - 1)}{x})").number_format = MONEY
    _widths(wc, [28, 28, 26] + [12] * len(meses_c) + [14, 10])

    # Cuenta declarada en Talana vs concepto real leído en los adjuntos
    wv = wb.create_sheet("Declarado vs Detectado")
    R_DECL = f"{D}!$E$2:$E${last}"
    decls = sorted({res["rendicion"]["concepto_declarado"] for res in results if res["rendicion"]["concepto_declarado"]})
    _header(wv, ["Cuenta declarada (Talana)"] + used + ["Total", "% identificado"])
    tv = len(used) + 2
    for i, dcl in enumerate(decls, 2):
        wv.cell(row=i, column=1, value=dcl)
        for j in range(2, tv):
            col = get_column_letter(j)
            wv.cell(row=i, column=j, value=f"=SUMIFS({R_MONTO},{R_DECL},$A{i},{R_CONC},{col}$1)").number_format = MONEY
        wv.cell(row=i, column=tv, value=f"=SUM(B{i}:{get_column_letter(tv - 1)}{i})").number_format = MONEY
    tr_v = len(decls) + 2
    _total_row(wv, tr_v, 2, tv)
    sin_col = get_column_letter(used.index(SIN_ID) + 2) if SIN_ID in used else None
    tot_col = get_column_letter(tv)
    for i in range(2, tr_v + 1):
        f = f"=IF({tot_col}{i}=0,\"\",1-{sin_col}{i}/{tot_col}{i})" if sin_col else f"=IF({tot_col}{i}=0,\"\",1)"
        wv.cell(row=i, column=tv + 1, value=f).number_format = "0.0%"
    wv.freeze_panes = "B2"
    _widths(wv, [34] + [14] * (len(used) + 2))
    wv.cell(row=tr_v + 2, column=1, value=(
        "Lectura: cada fila muestra en qué se gastó realmente lo rendido bajo esa cuenta. "
        "SIN_IDENTIFICAR = monto declarado que no se pudo leer en los adjuntos (sin adjunto, ilegible, "
        "falta OCR/7-Zip). % identificado = parte del monto explicada por los respaldos.")).font = Font(italic=True, size=9)

    wm = wb.create_sheet("Concepto x Mes")
    meses = sorted({parse_date(res["rendicion"].get("fecha")).strftime("%Y-%m")
                    for res in results if parse_date(res["rendicion"].get("fecha"))})
    _header(wm, ["Concepto detectado"] + meses + ["Total"])
    tcm = len(meses) + 2
    for i, c in enumerate(used, 2):
        wm.cell(row=i, column=1, value=c)
        for j in range(2, tcm):
            col = get_column_letter(j)
            wm.cell(row=i, column=j, value=f"=SUMIFS({R_MONTO},{R_CONC},$A{i},{R_MES},{col}$1)").number_format = MONEY
        wm.cell(row=i, column=tcm, value=f"=SUM(B{i}:{get_column_letter(tcm - 1)}{i})").number_format = MONEY
    _total_row(wm, len(used) + 2, 2, tcm)
    _widths(wm, [24] + [13] * (len(meses) + 1))

    we = wb.create_sheet("Errores")
    _header(we, ["ID", "Rendidor", "Estado", "Detalle", "Adjunto"])
    k = 1
    for res in results:
        for e in res["errores"] or ([] if res["estado"] == "OK" else [res["estado"]]):
            k += 1
            r = res["rendicion"]
            for c, v in enumerate([r["id"], r["rendidor"], res["estado"], e, r.get("adjunto")], 1):
                we.cell(row=k, column=c, value=v)
    _widths(we, [11, 28, 20, 80, 60])

    wpar = wb.create_sheet("Parámetros")
    for i, (kk, vv) in enumerate(meta.items(), 1):
        wpar.cell(row=i, column=1, value=kk).font = Font(bold=True)
        wpar.cell(row=i, column=2, value=str(vv))
    _widths(wpar, [28, 90])

    wb.save(out_path)
    return out_path


# =====================================================================
# 11. EJECUCIÓN COMPLETA
# =====================================================================
def run_rpa(xlsx: Path, hoja: str, motor: str, api_key: str = "", solo_movilizacion=False, rendidor=None,
            limite=None, reprocesar=False, progress=None, stop_event: threading.Event | None = None, cecos: list[str] | None = None, solo_finalizadas: bool = False,
            solo_descargar: bool = False) -> Path:
    rendiciones = read_rendiciones(xlsx, hoja, solo_movilizacion, rendidor, limite, cecos, solo_finalizadas)
    if not rendiciones:
        raise ValueError("No hay rendiciones que cumplan los filtros.")
    if cecos:
        log.info("CECOs seleccionados (%d): %s", len(cecos), ", ".join(cecos))
    if solo_descargar:
        dest = xlsx.parent / f"Adjuntos_por_CECO_{dt.datetime.now():%Y%m%d_%H%M}"
        return download_only(rendiciones, dest, xlsx.parent / f"{xlsx.stem}_rpa_cache", progress, stop_event)
    clf = Classifier(CONFIG["conceptos"])
    if motor == "claude":
        if not api_key:
            raise ValueError("El motor Claude requiere API key.")
        analyzer = ClaudeAnalyzer(clf, api_key)
        log.info("Motor: Claude (%s)", CONFIG["claude"]["modelo"])
    else:
        analyzer = LocalAnalyzer(clf)
        log.info("Motor: Local (texto + OCR Tesseract + reglas)")

    cache = xlsx.parent / f"{xlsx.stem}_rpa_cache"
    results: list = [None] * len(rendiciones)
    total, done = len(rendiciones), 0
    with ThreadPoolExecutor(max_workers=CONFIG["descarga"]["workers"]) as ex:
        futs = {ex.submit(process_one, r, analyzer, cache, reprocesar): i for i, r in enumerate(rendiciones)}
        for fut in as_completed(futs):
            i = futs[fut]
            r = rendiciones[i]
            try:
                results[i] = fut.result()
            except Exception as e:  # noqa: BLE001
                log.exception("Fallo inesperado en ID %s", r.id)
                results[i] = {"rendicion": asdict(r), "estado": "ERROR", "archivos": [], "items": [], "errores": [str(e)]}
            done += 1
            res = results[i]
            log.info("[%d/%d] %s | %s | %s | %d ítems", done, total, r.id, r.rendidor[:25], res["estado"], len(res["items"]))
            if progress:
                progress(done, total)
            if stop_event is not None and stop_event.is_set():
                log.warning("Detenido por el usuario: se generará el reporte con lo procesado.")
                for f in futs:
                    f.cancel()
                break
    results = [r for r in results if r is not None]
    audit(results)
    add_cuadratura(results)

    out = xlsx.parent / f"Desglose_Rendiciones_{dt.datetime.now():%Y%m%d_%H%M}.xlsx"
    meta = {"Archivo de entrada": xlsx.resolve(), "Hoja": hoja,
            "Motor": analyzer.name + (f" ({CONFIG['claude']['modelo']})" if analyzer.name == "claude" else ""),
            "Filtros": f"solo_movilizacion={solo_movilizacion}, solo_finalizadas={solo_finalizadas}, "
                       f"rendidor={rendidor}, limite={limite}",
            "CECOs": ", ".join(cecos) if cecos else "Todos",
            "Fecha ejecución": dt.datetime.now().strftime("%d-%m-%Y %H:%M"),
            "Rendiciones procesadas": len(results),
            "Con alertas": sum(1 for r in results if r.get("alertas")),
            "Con problemas de lectura": sum(1 for r in results if r["estado"] != "OK"),
            "Caché (re-ejecutar no vuelve a descargar)": cache}
    write_report(results, out, meta)
    for k in ("Rendiciones procesadas", "Con alertas", "Con problemas de lectura"):
        log.info("%-26s %s", k, meta[k])
    log.info("Reporte generado: %s", out)
    return out


# =====================================================================
# 12. INTERFAZ GRÁFICA
# =====================================================================
def check_tools(motor: str) -> list[str]:
    faltan = []
    if motor == "local":
        cmd = find_tesseract()
        if not cmd:
            faltan.append("Tesseract OCR (fotos y PDF escaneados): github.com/UB-Mannheim/tesseract/wiki")
        else:
            try:
                import pytesseract
                pytesseract.pytesseract.tesseract_cmd = cmd
                pytesseract.get_tesseract_version()
            except Exception:  # noqa: BLE001
                faltan.append(f"Tesseract encontrado en {cmd} pero no arrancó (¿instalación incompleta?)")
    if not find_7zip():
        faltan.append("7-Zip o UnRAR (archivos .rar): www.7-zip.org")
    return faltan


class QueueHandler(logging.Handler):
    def __init__(self, q):
        super().__init__()
        self.q = q

    def emit(self, record):
        self.q.put(("log", self.format(record)))


def launch_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title("RPA Rendiciones — Desglose por concepto")
    root.geometry("860x620")
    q: queue.Queue = queue.Queue()
    h = QueueHandler(q)
    h.setFormatter(logging.Formatter("%(asctime)s  %(message)s", "%H:%M:%S"))
    log.addHandler(h)
    stop_event = threading.Event()

    v_file, v_sheet = tk.StringVar(), tk.StringVar()
    v_motor = tk.StringVar(value="claude" if os.getenv("ANTHROPIC_API_KEY") else "local")
    v_key = tk.StringVar(value=os.getenv("ANTHROPIC_API_KEY", ""))
    v_mov, v_rep = tk.BooleanVar(), tk.BooleanVar()
    v_fin, v_dl = tk.BooleanVar(value=True), tk.BooleanVar()
    v_lim, v_rend = tk.StringVar(), tk.StringVar()
    state = {"out": None}

    frm = ttk.Frame(root, padding=12)
    frm.pack(fill="both", expand=True)
    frm.columnconfigure(1, weight=1)

    def pick_file():
        p = filedialog.askopenfilename(title="Selecciona el Excel de rendiciones",
                                       filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")])
        if not p:
            return
        v_file.set(p)
        try:
            sheets = list_sheets(Path(p))
            cb_sheet["values"] = sheets
            v_sheet.set(sheets[0])
            load_cecos()
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Error", f"No pude abrir el Excel:\n{e}")

    def load_cecos(*_):
        lb_ceco.delete(0, "end")
        if not v_file.get() or not v_sheet.get():
            return
        try:
            for c in list_cecos(Path(v_file.get()), v_sheet.get()):
                lb_ceco.insert("end", c)
        except Exception as e:  # noqa: BLE001
            log.warning("No pude leer los CECOs: %s", e)

    ttk.Label(frm, text="1. Excel de entrada").grid(row=0, column=0, sticky="w")
    ttk.Entry(frm, textvariable=v_file).grid(row=0, column=1, sticky="ew", padx=6)
    ttk.Button(frm, text="Seleccionar...", command=pick_file).grid(row=0, column=2)

    ttk.Label(frm, text="2. Hoja").grid(row=1, column=0, sticky="w", pady=4)
    cb_sheet = ttk.Combobox(frm, textvariable=v_sheet, state="readonly")
    cb_sheet.grid(row=1, column=1, sticky="ew", padx=6)
    cb_sheet.bind("<<ComboboxSelected>>", load_cecos)

    ttk.Label(frm, text="3. Motor de lectura").grid(row=2, column=0, sticky="w")
    mf = ttk.Frame(frm)
    mf.grid(row=2, column=1, sticky="w", padx=6)
    ttk.Radiobutton(mf, text="Local (gratis, OCR Tesseract)", variable=v_motor, value="local",
                   command=lambda: refresh_tools()).pack(side="left")
    ttk.Radiobutton(mf, text="Claude (mejor en fotos/escaneos)", variable=v_motor, value="claude").pack(side="left", padx=10)

    ttk.Label(frm, text="   API key Anthropic").grid(row=3, column=0, sticky="w", pady=4)
    ttk.Entry(frm, textvariable=v_key, show="•").grid(row=3, column=1, sticky="ew", padx=6)

    v_tess, v_7z = tk.StringVar(), tk.StringVar()

    def refresh_tools():
        v_tess.set(find_tesseract() or "no encontrado")
        v_7z.set(find_7zip() or "no encontrado")

    def pick_tool(titulo: str, setting_key: str):
        p = filedialog.askopenfilename(title=titulo, filetypes=[("Ejecutable", "*.exe"), ("Todos", "*.*")])
        if p:
            save_setting(setting_key, p)
            refresh_tools()

    tf = ttk.LabelFrame(frm, text="Programas detectados (sólo motor Local)", padding=6)
    tf.grid(row=4, column=0, columnspan=3, sticky="ew", pady=6)
    tf.columnconfigure(1, weight=1)
    ttk.Label(tf, text="Tesseract:").grid(row=0, column=0, sticky="w")
    ttk.Label(tf, textvariable=v_tess, foreground="#555").grid(row=0, column=1, sticky="w", padx=6)
    ttk.Button(tf, text="Buscar...", command=lambda: pick_tool(
        "Selecciona tesseract.exe", "tesseract_cmd")).grid(row=0, column=2)
    ttk.Label(tf, text="7-Zip / UnRAR:").grid(row=1, column=0, sticky="w", pady=(4, 0))
    ttk.Label(tf, textvariable=v_7z, foreground="#555").grid(row=1, column=1, sticky="w", padx=6, pady=(4, 0))
    ttk.Button(tf, text="Buscar...", command=lambda: pick_tool(
        "Selecciona 7z.exe o UnRAR.exe", "sevenzip_cmd")).grid(row=1, column=2, pady=(4, 0))
    refresh_tools()

    of = ttk.LabelFrame(frm, text="Filtros (opcional)", padding=6)
    of.grid(row=5, column=0, columnspan=3, sticky="ew", pady=6)
    ttk.Checkbutton(of, text="Sólo Movilización", variable=v_mov).grid(row=0, column=0, sticky="w")  # noqa: keep-row0
    ttk.Label(of, text="Rendidor contiene:").grid(row=0, column=1, padx=(16, 4))
    ttk.Entry(of, textvariable=v_rend, width=20).grid(row=0, column=2)
    ttk.Label(of, text="Límite (prueba):").grid(row=0, column=3, padx=(16, 4))
    ttk.Entry(of, textvariable=v_lim, width=6).grid(row=0, column=4)
    ttk.Checkbutton(of, text="Reprocesar (ignorar caché)", variable=v_rep).grid(row=0, column=5, padx=(16, 0))
    ttk.Checkbutton(of, text="Sólo pagadas / finalizadas", variable=v_fin).grid(row=1, column=0, sticky="w", pady=(6, 0))
    ttk.Checkbutton(of, text="Sólo descargar y ordenar en carpetas CECO/Rendidor (para revisar con Cowork)",
                    variable=v_dl).grid(row=1, column=1, columnspan=5, sticky="w", pady=(6, 0))
    ttk.Label(of, text="CECOs (Ctrl/Shift+clic; ninguno = todos):").grid(row=2, column=0, columnspan=2, sticky="nw", pady=(6, 0))
    lf = ttk.Frame(of)
    lf.grid(row=2, column=2, columnspan=4, sticky="ew", pady=(6, 0))
    lb_ceco = tk.Listbox(lf, selectmode="extended", height=6, exportselection=False)
    lb_ceco.pack(side="left", fill="x", expand=True)
    sb_c = ttk.Scrollbar(lf, command=lb_ceco.yview)
    sb_c.pack(side="left", fill="y")
    lb_ceco["yscrollcommand"] = sb_c.set

    def select_operacionales():
        objetivo = {normalize(c) for c in CONFIG["cecos_operacionales"]}
        lb_ceco.selection_clear(0, "end")
        n = 0
        for i in range(lb_ceco.size()):
            if normalize(lb_ceco.get(i)) in objetivo:
                lb_ceco.selection_set(i)
                n += 1
        log.info("CECOs operacionales seleccionados: %d de %d", n, len(objetivo))

    bc = ttk.Frame(of)
    bc.grid(row=3, column=2, columnspan=4, sticky="w", pady=(4, 0))
    ttk.Button(bc, text="CECOs operacionales (14)", command=select_operacionales).pack(side="left")
    ttk.Button(bc, text="Limpiar", command=lambda: lb_ceco.selection_clear(0, "end")).pack(side="left", padx=6)

    bf = ttk.Frame(frm)
    bf.grid(row=6, column=0, columnspan=3, sticky="ew", pady=4)
    btn_run = ttk.Button(bf, text="▶  Iniciar")
    btn_run.pack(side="left")
    btn_stop = ttk.Button(bf, text="■  Detener", state="disabled", command=lambda: stop_event.set())
    btn_stop.pack(side="left", padx=6)
    btn_open = ttk.Button(bf, text="Abrir reporte", state="disabled",
                          command=lambda: os.startfile(state["out"]) if hasattr(os, "startfile") else None)
    btn_open.pack(side="left")
    pb = ttk.Progressbar(bf, mode="determinate")
    pb.pack(side="left", fill="x", expand=True, padx=10)
    lbl = ttk.Label(bf, text="")
    lbl.pack(side="left")

    txt = tk.Text(frm, height=20, wrap="none", font=("Consolas", 9))
    txt.grid(row=7, column=0, columnspan=3, sticky="nsew")
    frm.rowconfigure(7, weight=1)
    sb = ttk.Scrollbar(frm, command=txt.yview)
    sb.grid(row=7, column=3, sticky="ns")
    txt["yscrollcommand"] = sb.set

    def start():
        if not v_file.get() or not Path(v_file.get()).exists():
            messagebox.showwarning("Falta archivo", "Selecciona el Excel de rendiciones.")
            return
        if v_motor.get() == "claude" and not v_key.get().strip() and not v_dl.get():
            messagebox.showwarning("Falta API key", "El motor Claude necesita una API key de Anthropic.")
            return
        try:
            limite = int(v_lim.get()) if v_lim.get().strip() else None
        except ValueError:
            messagebox.showwarning("Límite", "El límite debe ser un número.")
            return
        faltan = check_tools("descarga" if v_dl.get() else v_motor.get())
        if faltan and not messagebox.askyesno(
                "Faltan programas", "No encontré:\n\n• " + "\n• ".join(faltan) +
                "\n\nLo que dependa de ellos quedará como error (y se reintentará en la próxima corrida "
                "cuando los instales).\n\n¿Continuar igual?"):
            return
        stop_event.clear()
        btn_run["state"], btn_stop["state"], btn_open["state"] = "disabled", "normal", "disabled"
        pb["value"] = 0

        sel_cecos = [lb_ceco.get(i) for i in lb_ceco.curselection()]

        def work():
            try:
                out = run_rpa(Path(v_file.get()), v_sheet.get(), v_motor.get(), v_key.get().strip(), v_mov.get(),
                              v_rend.get().strip() or None, limite, v_rep.get(),
                              progress=lambda d, t: q.put(("prog", (d, t))), stop_event=stop_event,
                              cecos=sel_cecos or None, solo_finalizadas=v_fin.get(), solo_descargar=v_dl.get())
                q.put(("done", out))
            except Exception as e:  # noqa: BLE001
                log.exception("Error")
                q.put(("error", str(e)))

        threading.Thread(target=work, daemon=True).start()

    btn_run["command"] = start

    def poll():
        try:
            while True:
                kind, payload = q.get_nowait()
                if kind == "log":
                    txt.insert("end", payload + "\n")
                    txt.see("end")
                elif kind == "prog":
                    d, t = payload
                    pb["maximum"], pb["value"] = t, d
                    lbl["text"] = f"{d}/{t}"
                elif kind in ("done", "error"):
                    btn_run["state"], btn_stop["state"] = "normal", "disabled"
                    if kind == "done":
                        state["out"] = str(payload)
                        btn_open["state"] = "normal"
                        messagebox.showinfo("Listo", f"Reporte generado:\n{payload}")
                    else:
                        messagebox.showerror("Error", payload)
        except queue.Empty:
            pass
        root.after(150, poll)

    poll()
    root.mainloop()


# =====================================================================
# 13. MAIN
# =====================================================================
def main():
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)-7s | %(message)s")
    for noisy in ("httpx", "urllib3", "anthropic", "PIL"):
        logging.getLogger(noisy).setLevel(logging.WARNING)

    if "--cli" not in sys.argv:
        try:
            launch_gui()
            return
        except ImportError:
            print("Tkinter no disponible: usa el modo --cli")
            return

    p = argparse.ArgumentParser()
    p.add_argument("--cli", required=True, metavar="EXCEL")
    p.add_argument("--hoja", default=CONFIG["hojas_soportadas"][0])
    p.add_argument("--motor", choices=["local", "claude"], default="local")
    p.add_argument("--solo-movilizacion", action="store_true")
    p.add_argument("--rendidor")
    p.add_argument("--limite", type=int)
    p.add_argument("--reprocesar", action="store_true")
    p.add_argument("--ceco", action="append", help="Repetible: --ceco \"230 - SANTIAGO\" --ceco \"290 - TRADE\"")
    p.add_argument("--solo-finalizadas", action="store_true")
    p.add_argument("--cecos-operacionales", action="store_true", help="Usa los 14 CECOs de CONFIG['cecos_operacionales']")
    p.add_argument("--solo-descargar", action="store_true", help="Sólo descarga y ordena en carpetas CECO/Rendidor")
    a = p.parse_args()
    run_rpa(Path(a.cli), a.hoja, a.motor, os.getenv("ANTHROPIC_API_KEY", ""), a.solo_movilizacion,
            a.rendidor, a.limite, a.reprocesar,
            cecos=(CONFIG["cecos_operacionales"] if a.cecos_operacionales else a.ceco),
            solo_finalizadas=a.solo_finalizadas, solo_descargar=a.solo_descargar)


if __name__ == "__main__":
    main()
