# -*- coding: utf-8 -*-
"""Consolida en un solo Excel: el archivo original + todas las hojas del desglose revisado + un Resumen Ejecutivo."""
import glob
from copy import copy

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

ORIG = "Presupuestos_Rendiciones_actualizado.xlsx"
DES = sorted(glob.glob("Desglose_Rendiciones_*_revisado.xlsx"))[-1]
OUT = "Presupuestos_Rendiciones_CONSOLIDADO.xlsx"
RENOMBRE = {"Rendiciones CECO-Rendidor": "CECO-Rendidor Detallado"}

wb = openpyxl.load_workbook(ORIG)
src = openpyxl.load_workbook(DES)
nuevas = []
for ws in src.worksheets:
    name = RENOMBRE.get(ws.title, ws.title)
    dst = wb.create_sheet(name)
    nuevas.append(name)
    for row in ws.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font, d.fill, d.border = copy(c.font), copy(c.fill), copy(c.border)
                d.alignment, d.number_format = copy(c.alignment), c.number_format
            if c.hyperlink:
                d.hyperlink = c.hyperlink.target
    for k, dim in ws.column_dimensions.items():
        dst.column_dimensions[k].width = dim.width
    for k, dim in ws.row_dimensions.items():
        if dim.outlineLevel:
            dst.row_dimensions[k].outlineLevel = dim.outlineLevel
        if dim.height:
            dst.row_dimensions[k].height = dim.height
    dst.freeze_panes = ws.freeze_panes
    if ws.auto_filter.ref:
        dst.auto_filter.ref = ws.auto_filter.ref
    dst.sheet_properties.outlinePr.summaryBelow = ws.sheet_properties.outlinePr.summaryBelow
    print("copiada", name, ws.max_row)

# ---- Resumen Ejecutivo ----
b = pd.read_excel(DES, sheet_name="Base Dinámica")
we = wb.create_sheet("Resumen Ejecutivo", 0)
H = PatternFill("solid", start_color="1F4E78")
MONEY = '_ * #,##0_ ;_ * \\-#,##0_ ;_ * "-"_ ;_ @_ '
r = 1
we.cell(row=r, column=1, value="Rendiciones 2026 — CECOs operacionales: gasto real leído en los adjuntos").font = Font(bold=True, size=14)
r += 1
we.cell(row=r, column=1, value=f"Fuente: Talana (Reembolsos Consolidado, finalizadas) + lectura de adjuntos. Detalle: {DES}").font = Font(italic=True, size=9)
r += 2


def tabla(titulo, df, r, pct_col=True):
    we.cell(row=r, column=1, value=titulo).font = Font(bold=True, size=11)
    r += 1
    cols = list(df.columns)
    for j, h in enumerate([df.index.name or ""] + cols, 1):
        x = we.cell(row=r, column=j, value=h)
        x.fill, x.font = H, Font(bold=True, color="FFFFFF")
        x.alignment = Alignment(wrap_text=True)
    for i, (k, vals) in enumerate(df.iterrows(), r + 1):
        we.cell(row=i, column=1, value=str(k))
        for j, (col, v) in enumerate(zip(cols, vals), 2):
            x = we.cell(row=i, column=j, value=float(v) if v == v else None)
            x.number_format = "0.0%" if str(col).startswith("%") else MONEY
    return r + len(df) + 3


tot = b["Monto"].sum()
gc = b.groupby("Concepto")["Monto"].sum().sort_values(ascending=False).to_frame("Monto")
gc["% del total"] = gc["Monto"] / tot
gc.index.name = "Concepto"
r = tabla(f"1. Gasto por concepto (total {tot:,.0f})".replace(",", "."), gc, r)
b["Nivel de detalle"] = "Concepto identificado"
b.loc[b["Concepto"] == "OTROS", "Nivel de detalle"] = "OTROS con subtipo"
b.loc[b["Concepto probable (sin detalle)"].notna(), "Nivel de detalle"] = "Sin detalle, con concepto probable"
b.loc[(b["Concepto"] == "SIN_IDENTIFICAR") & b["Concepto probable (sin detalle)"].isna(), "Nivel de detalle"] = \
    "Sin detalle ni probable (ver Pendientes)"
gn = b.groupby("Nivel de detalle")["Monto"].sum().to_frame("Monto")
gn["% del total"] = gn["Monto"] / tot
r = tabla("2. Calidad del detalle", gn, r)
go = b[b["Concepto"] == "OTROS"].groupby("Subtipo OTROS")["Monto"].sum().sort_values(ascending=False).to_frame("Monto")
go["% de OTROS"] = go["Monto"] / go["Monto"].sum()
r = tabla("3. Qué hay dentro de OTROS", go, r)
gp = b[b["Concepto probable (sin detalle)"].notna()].groupby("Concepto probable (sin detalle)")["Monto"].sum() \
    .sort_values(ascending=False).to_frame("Monto")
r = tabla("4. Monto sin detalle según concepto probable", gp, r)
gce = b.pivot_table(index="CECO", columns="Concepto", values="Monto", aggfunc="sum").fillna(0)
gce = gce[gc.index]
gce.insert(0, "Total", gce.sum(axis=1))
r = tabla("5. CECO x Concepto", gce, r)
we.cell(row=r, column=1, value="Hojas agregadas").font = Font(bold=True, size=11)
r += 1
guia = {
    "CECO-Rendidor Detallado": "Evolutivo mensual CECO → Rendidor → Concepto detallado (agrupable, con totales y %)",
    "Top 5 Conceptos": "5 conceptos con más gasto por CECO y por rendidor",
    "Cabify + Mov Evolutivo": "Usuarios Cabify que además rinden apps o combustible, mes a mes",
    "Desglose OTROS": "OTROS por subtipo x CECO, x mes, x rendidor y principales emisores",
    "Pendientes de detalle": "Rendiciones con monto sin detalle, con probable, motivo y link al adjunto (para control)",
    "Desglose x Ítem": "Cada comprobante/ítem leído: emisor, fecha, monto, concepto, método, alertas, observación",
    "Base Dinámica": "Tabla plana para tablas dinámicas (incluye concepto detallado y probable)",
    "Parámetros": "Método, filtros y métricas de la lectura",
}
for k, v in guia.items():
    we.cell(row=r, column=1, value=k).font = Font(bold=True)
    we.cell(row=r, column=2, value=v)
    r += 1
we.column_dimensions["A"].width = 55
for col in "BCDEFGHIJKLMNOPQRSTU":
    we.column_dimensions[col].width = 14

# orden: Resumen Ejecutivo + hojas nuevas principales primero, luego las originales
orden = ["Resumen Ejecutivo", "CECO-Rendidor Detallado", "Top 5 Conceptos", "Cabify + Mov Evolutivo",
         "Desglose OTROS", "Pendientes de detalle"]
wb._sheets = [wb[n] for n in orden] + [s for s in wb._sheets if s.title not in orden]
wb.active = 0
wb.save(OUT)
print("OK ->", OUT, "| hojas:", len(wb.sheetnames))
