#!/usr/bin/env python3
"""Asistente interactivo para generar la base única de Power BI.

Te pregunta (con ventanas de selección o, si no hay entorno gráfico, por consola):
  1. Dónde están los archivos: un Excel de Talana o una carpeta con varios.
  2. Qué hojas analizar: solicitudes de Talana (obligatoria), extracto Cabify y cruce Talana↔Cabify (opcionales).
  3. En qué carpeta guardar la base.
Luego ejecuta generar_base_powerbi.py con esas elecciones. Recuerda la última selección.

Uso:   python generar_base_interactivo.py          (requiere los demás .py del proyecto en la misma carpeta)
       python generar_base_interactivo.py --consola   (fuerza el modo consola)
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent
PREFS = AQUI / ".ultima_seleccion.json"
SUGERIDAS = {"rend": ("reembolsos consolidado", "reembolso", "rendicion", "talana"),
             "cab": ("cabify extracto", "cabify"),
             "aud": ("auditoria cabify", "auditoria")}


# ---------------------------------------------------------------- entrada/salida
class UI:
    def __init__(self, consola: bool):
        self.tk = None
        if not consola:
            try:
                import tkinter as tk
                from tkinter import filedialog
                root = tk.Tk()
                root.withdraw()
                root.attributes("-topmost", True)
                self.tk, self.fd, self.root = tk, filedialog, root
            except Exception:
                self.tk = None          # sin entorno gráfico → consola

    def carpeta(self, titulo: str, inicial: str = "") -> str:
        if self.tk:
            r = self.fd.askdirectory(title=titulo, initialdir=inicial or str(Path.home()))
            return r or ""
        return input(f"{titulo}\n  ruta de la carpeta [{inicial}]: ").strip().strip('"') or inicial

    def archivo(self, titulo: str, inicial: str = "") -> str:
        if self.tk:
            r = self.fd.askopenfilename(title=titulo, initialdir=inicial or str(Path.home()),
                                        filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")])
            return r or ""
        return input(f"{titulo}\n  ruta del archivo [{inicial}]: ").strip().strip('"') or inicial

    def elegir(self, titulo: str, opciones: list[str], multiple=False, defecto: list[int] | None = None,
               opcional=False) -> list[int]:
        """Menú numerado en consola (también en modo gráfico: es más simple y portable). Devuelve índices."""
        print(f"\n{titulo}")
        for i, o in enumerate(opciones, 1):
            marca = " *" if defecto and (i - 1) in defecto else ""
            print(f"  {i:>2}) {o}{marca}")
        ayuda = "números separados por coma" if multiple else "un número"
        if opcional:
            ayuda += ", 0 = ninguna"
        pref = ",".join(str(i + 1) for i in (defecto or []))
        while True:
            r = input(f"  Elige ({ayuda}) [{pref}]: ").strip() or pref
            if opcional and r in ("0", "-"):
                return []
            try:
                idx = [int(x) - 1 for x in r.replace(" ", "").split(",") if x]
                if idx and all(0 <= i < len(opciones) for i in idx) and (multiple or len(idx) == 1):
                    return idx
            except ValueError:
                pass
            print("  Opción no válida.")

    def si_no(self, pregunta: str, defecto=False) -> bool:
        r = input(f"{pregunta} [{'S/n' if defecto else 's/N'}]: ").strip().lower()
        return defecto if not r else r in ("s", "si", "sí", "y", "yes")


def hojas_de(xlsx: Path) -> list[tuple[str, str]]:
    """[(nombre, 'filas x columnas')] sin cargar el libro completo."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out = [(ws.title, f"{ws.max_row or '?'} filas × {ws.max_column or '?'} cols") for ws in wb.worksheets]
    wb.close()
    return out


def sugerir(nombres: list[str], clave: str) -> list[int]:
    for pat in SUGERIDAS[clave]:
        for i, n in enumerate(nombres):
            if pat in n.lower() and (clave != "cab" or "audit" not in n.lower()):
                return [i]
    return []


def cargar_prefs() -> dict:
    try:
        return json.loads(PREFS.read_text(encoding="utf-8"))
    except Exception:
        return {}


# ---------------------------------------------------------------- flujo
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--consola", action="store_true", help="no usar ventanas de selección")
    a = ap.parse_args()
    prefs = cargar_prefs()
    ui = UI(a.consola)

    print("=== Base única de rendiciones para Power BI ===")

    # 1) origen
    modo = ui.elegir("¿Dónde están los datos?", ["Un archivo Excel", "Una carpeta con varios Excel"], defecto=[0])[0]
    if modo == 0:
        ruta = ui.archivo("Selecciona el Excel de Talana", prefs.get("carpeta_origen", ""))
        xlsx = Path(ruta) if ruta else None
    else:
        carpeta = ui.carpeta("Selecciona la carpeta con los Excel", prefs.get("carpeta_origen", ""))
        if not carpeta:
            sys.exit("No se seleccionó carpeta.")
        libros = sorted(p for p in Path(carpeta).glob("*.xls[xm]") if not p.name.startswith("~$"))
        if not libros:
            sys.exit(f"No hay archivos .xlsx/.xlsm en {carpeta}")
        xlsx = libros[ui.elegir("Archivos encontrados:", [p.name for p in libros])[0]]
    if not xlsx or not xlsx.exists():
        sys.exit("No se seleccionó un archivo válido.")
    print(f"\nArchivo: {xlsx}")

    # 2) hojas
    info = hojas_de(xlsx)
    nombres = [n for n, _ in info]
    etiquetas = [f"{n}   ({d})" for n, d in info]
    i_rend = ui.elegir("HOJA A ANALIZAR — solicitudes/rendiciones de Talana (obligatoria):", etiquetas,
                       defecto=sugerir(nombres, "rend"))[0]
    resto = [i for i in range(len(nombres)) if i != i_rend]
    sel_cab = ui.elegir("HOJA con el extracto de viajes Cabify (opcional; 0 = no incluir Cabify):",
                        [etiquetas[i] for i in resto], defecto=[resto.index(i) for i in sugerir(nombres, "cab") if i in resto],
                        opcional=True)
    i_cab = resto[sel_cab[0]] if sel_cab else None
    i_aud = None
    if i_cab is not None:
        resto2 = [i for i in resto if i != i_cab]
        sel_aud = ui.elegir("HOJA con el cruce manual Talana↔Cabify (opcional, mejora el cruce de nombres):",
                            [etiquetas[i] for i in resto2],
                            defecto=[resto2.index(i) for i in sugerir(nombres, "aud") if i in resto2], opcional=True)
        i_aud = resto2[sel_aud[0]] if sel_aud else None

    # 3) destino
    destino = ui.carpeta("Selecciona la carpeta donde guardar la base",
                         prefs.get("carpeta_destino", str(xlsx.parent / "salida_powerbi")))
    if not destino:
        sys.exit("No se seleccionó carpeta de destino.")
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)

    nombre = input(f"\nNombre del archivo de salida [{prefs.get('nombre', 'Base_Rendiciones_PowerBI')}]: ").strip() \
        or prefs.get("nombre", "Base_Rendiciones_PowerBI")
    reprocesar = ui.si_no("¿Volver a leer los adjuntos desde cero (más lento)?", False)

    # 4) resumen y confirmación
    print("\n--- Resumen ---")
    print(f"  Excel         : {xlsx}")
    print(f"  Hoja Talana   : {nombres[i_rend]}")
    print(f"  Hoja Cabify   : {nombres[i_cab] if i_cab is not None else '(no se incluye)'}")
    if i_cab is not None:
        print(f"  Hoja cruce    : {nombres[i_aud] if i_aud is not None else '(ninguna)'}")
    print(f"  Guardar en    : {destino}")
    print(f"  Nombre        : {nombre}")
    if not ui.si_no("¿Ejecutar?", True):
        sys.exit("Cancelado.")

    PREFS.write_text(json.dumps({"carpeta_origen": str(xlsx.parent), "carpeta_destino": str(destino),
                                 "nombre": nombre}, ensure_ascii=False), encoding="utf-8")

    # 5) ejecutar el generador. Los archivos de trabajo del archivo elegido quedan en el destino
    #    (así datos de un Excel no se mezclan con los de otro).
    sys.path.insert(0, str(AQUI))
    import generar_base_powerbi as G
    trabajo = destino / "_trabajo"
    trabajo.mkdir(exist_ok=True)
    argv = ["--excel", str(xlsx), "--salida", str(destino), "--nombre", nombre,
            "--hoja-rendiciones", nombres[i_rend], "--pkl", str(trabajo / f"{xlsx.stem}_resultados.pkl"),
            "--estado-pkl", str(trabajo / f"{xlsx.stem}_estado.pkl")]
    for opt, ruta in (("--lecturas", AQUI / "lecturas"), ("--indice", AQUI / "revisar" / "indice.json")):
        if ruta.exists():
            argv += [opt, str(ruta)]
    if i_cab is None:
        argv.append("--sin-cabify")
    else:
        argv += ["--hoja-cabify", nombres[i_cab]]
        if i_aud is not None:
            argv += ["--hoja-auditoria", nombres[i_aud]]
    if reprocesar:
        argv.append("--reprocesar")
        for f in trabajo.glob(f"{xlsx.stem}_*.pkl"):      # evita reutilizar estado viejo
            f.unlink()
    G.main(argv)
    print(f"\nListo. Archivos en: {destino}")


if __name__ == "__main__":
    main()
