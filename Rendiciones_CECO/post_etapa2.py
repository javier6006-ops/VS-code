# -*- coding: utf-8 -*-
"""Etapa 2 del postproceso: incorpora las lecturas manuales (lecturas/lote_*.json) a los resultados del RPA,
recalcula auditoría y cuadratura, genera el Desglose y agrega hojas de análisis:
  - Rendiciones CECO-Rendidor : evolutivo mensual CECO → Rendidor → Concepto (total y por concepto)
  - Top 5 Conceptos           : 5 conceptos con más gasto por CECO y por CECO-Rendidor
  - Cabify + Mov Evolutivo    : usuarios Cabify que además rinden APPS_TRANSPORTE o COMBUSTIBLE, mes a mes
  - Base Dinámica             : tabla plana para tablas dinámicas
Formato de cada lectura: {"id", "items": [{"img","emisor","fecha":"DD-MM-YYYY","monto","concepto","descripcion"}],
                          "ilegibles": ["03.jpg", ...], "nota": ""}"""
# Las reglas de lectura/clasificación viven en pipeline_rendiciones.py (compartido con generar_base_powerbi.py).
import datetime as dt
import json
import sys
import unicodedata
from pathlib import Path

sys.argv = ["x"]
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402

import pipeline_rendiciones as P  # noqa: E402
import rpa_rendiciones as R  # noqa: E402
from pipeline_rendiciones import GENERICOS, MESES_ES, motivo_sin, subtipo_otros  # noqa: E402

XLSX = Path("Presupuestos_Rendiciones_actualizado.xlsx")
CACHE = XLSX.parent / f"{XLSX.stem}_rpa_cache"

results = P.cargar_resultados(XLSX, Path("resultados_base.pkl"))
lecturas, index = P.cargar_lecturas(Path("lecturas"), Path("revisar/indice.json"))
estado = P.ejecutar_pipeline(results, lecturas, index, CACHE)
PROB = estado.prob

out = Path(f"Desglose_Rendiciones_{dt.datetime.now():%Y%m%d_%H%M}_revisado.xlsx")
meta = {"Archivo de entrada": XLSX.resolve(), "Hoja": "Reembolsos Consolidado", "Motor": "local + lectura manual de imágenes",
        "Filtros": "solo_finalizadas=True, CECOs operacionales (14)",
        "Fecha ejecución": dt.datetime.now().strftime("%d-%m-%Y %H:%M"), **estado.stats,
        "Nota montos": "Monto ítem = monto prorrateado a lo declarado; 'Monto bruto leído' = lo leído en el comprobante"}
R.write_report(results, out, meta)

# ---------- 2. Base plana ----------
def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(s.replace("-", " ").split())


rows = []
for res in results:
    r = res["rendicion"]
    f = R.parse_date(r.get("fecha"))
    for it in res["items"]:
        if it.get("monto"):
            rows.append({"CECO": r["ceco"], "Rendidor": r["rendidor"], "Usuario Cabify": r["usa_cabify"],
                         "Mes": f.strftime("%Y-%m") if f else "", "Concepto": it.get("concepto") or "OTROS",
                         "Subtipo OTROS": subtipo_otros(it),
                         "Concepto detallado": (f"OTROS · {subtipo_otros(it)}" if it.get("concepto") == "OTROS" else
                                                f"SIN_IDENTIFICAR · {motivo_sin(it)}" if it.get("concepto") == R.SIN_ID
                                                else it.get("concepto")),
                         "Concepto probable (sin detalle)": (PROB[r["id"]][0] if it.get("concepto") == R.SIN_ID or
                                                             subtipo_otros(it) in GENERICOS else ""),
                         "Base del probable": (PROB[r["id"]][1] if it.get("concepto") == R.SIN_ID or
                                               subtipo_otros(it) in GENERICOS else ""),
                         "Detalle de la rendición (carpetas/correos)": PROB[r["id"]][2],
                         "Emisor / detalle": (str(it.get("emisor") or it.get("descripcion") or "")[:80]
                                              if it.get("concepto") == "OTROS" else ""),
                         "Monto": float(it["monto"]), "ID": r["id"]})
df = pd.DataFrame(rows)
_p = df["Concepto probable (sin detalle)"].fillna("")
df.loc[_p != "", "Concepto detallado"] = df.loc[_p != "", "Concepto detallado"] + " → probable " + _p[_p != ""]
meses = sorted(m for m in df["Mes"].unique() if m)
nrend = {(r["rendicion"]["ceco"], r["rendicion"]["rendidor"]): 0 for r in results}
for res in results:
    nrend[(res["rendicion"]["ceco"], res["rendicion"]["rendidor"])] += 1

wb = openpyxl.load_workbook(out)
MONEY, HF, HFont = R.MONEY, R.HDR_FILL, R.HDR_FONT
wd = wb["Desglose x Ítem"]
for c, h in ((23, "Monto bruto leído"), (24, "Factor prorrateo"), (25, "Subtipo OTROS")):
    x = wd.cell(row=1, column=c, value=h)
    x.fill, x.font = HF, HFont
n = 1
for res in results:
    for it in res["items"]:
        n += 1
        wd.cell(row=n, column=23, value=it.get("monto_bruto")).number_format = MONEY
        wd.cell(row=n, column=24, value=it.get("factor")).number_format = "0.000"
        wd.cell(row=n, column=25, value=subtipo_otros(it))
wd.auto_filter.ref = f"A1:Y{max(n, 2)}"


def header(ws, hs, row=1):
    for c, h in enumerate(hs, 1):
        x = ws.cell(row=row, column=c, value=h)
        x.fill, x.font, x.border = HF, HFont, R.BORDER
        x.alignment = R.Alignment(horizontal="center", vertical="center", wrap_text=True)


def paint(ws, row, ncol, fill, bold=True):
    for c in range(1, ncol + 1):
        ws.cell(row=row, column=c).fill = fill
        if bold:
            ws.cell(row=row, column=c).font = Font(bold=True)


# ---------- 3. Rendiciones CECO-Rendidor (evolutivo) ----------
ws = wb.create_sheet("Rendiciones CECO-Rendidor", 1)
hs = ["CECO", "Rendidor", "Nivel", "Concepto"] + meses + ["Total", "Promedio mensual", "% del rendidor",
                                                          "% SIN_IDENTIFICAR", "N° rendiciones"]
header(ws, hs)
ws.freeze_panes = "E2"
m0, mT = 5, 5 + len(meses)          # columnas de meses y total
row = 2
ceco_rows = []
for ceco in sorted(df["CECO"].unique()):
    dc = df[df["CECO"] == ceco]
    ceco_row = row
    row += 1
    rend_rows = []
    for rend in dc.groupby("Rendidor")["Monto"].sum().sort_values(ascending=False).index:
        dr = dc[dc["Rendidor"] == rend]
        rr = row
        ws.cell(row=rr, column=1, value=ceco)
        ws.cell(row=rr, column=2, value=rend)
        ws.cell(row=rr, column=3, value="Rendidor")
        ws.cell(row=rr, column=4, value="Total mes")
        piv = dr.pivot_table(index="Concepto detallado", columns="Mes", values="Monto", aggfunc="sum").fillna(0)
        concs = piv.sum(axis=1).sort_values(ascending=False).index
        for j, m in enumerate(meses):
            ws.cell(row=rr, column=m0 + j, value=f"=SUM({L(m0 + j)}{rr + 1}:{L(m0 + j)}{rr + len(concs)})")
        ws.cell(row=rr, column=mT + 4, value=nrend.get((ceco, rend), 0))
        paint(ws, rr, len(hs), R.WARN_FILL)
        sin_rows = []
        for k, c in enumerate(concs, 1):
            x = rr + k
            ws.cell(row=x, column=1, value=ceco)
            ws.cell(row=x, column=2, value=rend)
            ws.cell(row=x, column=3, value="Concepto")
            ws.cell(row=x, column=4, value=c)
            for j, m in enumerate(meses):
                v = float(piv.loc[c, m]) if m in piv.columns else 0
                ws.cell(row=x, column=m0 + j, value=v or None)
            ws.cell(row=x, column=mT + 2, value=f"=IF({L(mT)}{rr}=0,\"\",{L(mT)}{x}/{L(mT)}{rr})")
            if str(c).startswith(R.SIN_ID):
                sin_rows.append(x)
                paint(ws, x, len(hs), R.PatternFill("solid", start_color="E7E6E6"), bold=False)
            ws.row_dimensions[x].outlineLevel = 2
        ws.cell(row=rr, column=mT + 3, value=(f"=IF({L(mT)}{rr}=0,\"\",(" + "+".join(f"{L(mT)}{x}" for x in sin_rows)
                                                 + f")/{L(mT)}{rr})") if sin_rows else 0)
        ws.row_dimensions[rr].outlineLevel = 1
        rend_rows.append(rr)
        row = rr + len(concs) + 1
    ws.cell(row=ceco_row, column=1, value=ceco)
    ws.cell(row=ceco_row, column=2, value=f"TOTAL {ceco}")
    ws.cell(row=ceco_row, column=3, value="CECO")
    ws.cell(row=ceco_row, column=4, value="Total mes")
    for j in range(len(meses)):
        col = L(m0 + j)
        ws.cell(row=ceco_row, column=m0 + j, value="=" + "+".join(f"{col}{x}" for x in rend_rows))
    ws.cell(row=ceco_row, column=mT + 4, value="=" + "+".join(f"{L(mT + 4)}{x}" for x in rend_rows))
    paint(ws, ceco_row, len(hs), R.TOTAL_FILL)
    ceco_rows.append(ceco_row)
ws.cell(row=row, column=1, value="TOTAL GENERAL")
ws.cell(row=row, column=3, value="Total")
for j in range(len(meses)):
    col = L(m0 + j)
    ws.cell(row=row, column=m0 + j, value="=" + "+".join(f"{col}{x}" for x in ceco_rows))
ws.cell(row=row, column=mT + 4, value="=" + "+".join(f"{L(mT + 4)}{x}" for x in ceco_rows))
paint(ws, row, len(hs), R.TOTAL_FILL)
for x in range(2, row + 1):
    ws.cell(row=x, column=mT, value=f"=SUM({L(m0)}{x}:{L(mT - 1)}{x})")
    ws.cell(row=x, column=mT + 1, value=f"=IFERROR(AVERAGEIF({L(m0)}{x}:{L(mT - 1)}{x},\">0\"),\"\")")
    for c in range(m0, mT + 2):
        ws.cell(row=x, column=c).number_format = MONEY
    for c in (mT + 2, mT + 3):
        ws.cell(row=x, column=c).number_format = "0.0%"
ws.auto_filter.ref = f"A1:{L(len(hs))}{row}"
ws.sheet_properties.outlinePr.summaryBelow = False
R._widths(ws, [26, 30, 10, 22] + [12] * len(meses) + [14, 13, 10, 11, 10])

# ---------- 4. Top 5 conceptos ----------
wt = wb.create_sheet("Top 5 Conceptos", 2)
header(wt, ["CECO", "Rendidor", "Ranking", "Concepto", "Monto", "% del total"])
wt.freeze_panes = "C2"
row, top_ceco = 2, {}
for ceco in sorted(df["CECO"].unique()):
    dc = df[df["CECO"] == ceco]
    for nivel, sub in [("(Todo el CECO)", dc)] + [(r, dc[dc["Rendidor"] == r]) for r in
                                                  dc.groupby("Rendidor")["Monto"].sum().sort_values(ascending=False).index]:
        s = sub.groupby("Concepto")["Monto"].sum().sort_values(ascending=False)
        tot = s.sum()
        if nivel == "(Todo el CECO)":
            top_ceco[ceco] = (tot, s.head(5))
        for k, (c, v) in enumerate(s.head(5).items(), 1):
            for col, val in enumerate([ceco, nivel, k, c, v, v / tot if tot else None], 1):
                wt.cell(row=row, column=col, value=val)
            wt.cell(row=row, column=5).number_format = MONEY
            wt.cell(row=row, column=6).number_format = "0.0%"
            if nivel == "(Todo el CECO)":
                paint(wt, row, 6, R.TOTAL_FILL)
            row += 1
wt.auto_filter.ref = f"A1:F{row - 1}"
R._widths(wt, [28, 32, 8, 24, 14, 10])

# ---------- 5. Cabify + Movilización evolutivo ----------
cab = pd.read_excel(XLSX, sheet_name="Cabify Extracto")
cab = cab[cab["Estado final"] != "No encontrado"].copy()
cab["Mes"] = cab.apply(lambda x: f"{int(x['Año'])}-{MESES_ES.index(str(x['Mes']).strip().capitalize()) + 1:02d}", axis=1)
cab["pn"] = cab["Pasajero"].map(norm)
aud = pd.read_excel(XLSX, sheet_name="Auditoria Cabify vs Rend", header=1)
manual_map = {}
for _, a in aud[aud["¿Coincide en Cabify?"] == "SI"].iterrows():
    manual_map[norm(a["Nombre (origen: Talana)"])] = {norm(n) for n in str(a["Nombre (origen: Cabify)"]).split(";")}
pas = set(cab["pn"])


def cabify_names(talana):
    t = norm(talana)
    if t in manual_map:
        return manual_map[t], "Auditoría Cabify vs Rend"
    tk = t.split()
    hits = {p for p in pas if len(p.split()) >= 2 and p.split()[0] == tk[0] and p.split()[-1] in tk[1:]}
    return hits, "Nombre + apellido" if hits else ""


mov = df[df["Concepto"].isin(["APPS_TRANSPORTE", "COMBUSTIBLE"])]
wcab = wb.create_sheet("Cabify + Mov Evolutivo", 3)
hs = ["CECO", "Rendidor (Talana)", "Pasajero(s) Cabify", "Match", "Fuente"] + meses + ["Total", "Promedio mensual"]
header(wcab, hs)
wcab.freeze_panes = "F2"
row = 2
f0 = 6
resumen_cab = []
for (ceco, rend), g in mov.groupby(["CECO", "Rendidor"]):
    names, how = cabify_names(rend)
    if not names:
        continue
    cg = cab[cab["pn"].isin(names) & cab["Mes"].isin(meses)]
    if cg["Precio Fnal"].sum() <= 0:                     # sin viajes Cabify en el período: no es usuario activo
        continue
    lines = [("Cabify corporativo (viajes)", cg.groupby("Mes")["Precio Fnal"].sum()),
             ("Rendido APPS_TRANSPORTE", g[g["Concepto"] == "APPS_TRANSPORTE"].groupby("Mes")["Monto"].sum()),
             ("Rendido COMBUSTIBLE", g[g["Concepto"] == "COMBUSTIBLE"].groupby("Mes")["Monto"].sum())]
    start = row
    for fuente, s in lines:
        for c, v in enumerate([ceco, rend, "; ".join(sorted(cg["Pasajero"].unique())), how, fuente], 1):
            wcab.cell(row=row, column=c, value=v)
        for j, m in enumerate(meses):
            wcab.cell(row=row, column=f0 + j, value=float(s.get(m, 0)) or None)
        row += 1
    for c, v in enumerate([ceco, rend, "", "", "Total movilización (Cabify + rendido)"], 1):
        wcab.cell(row=row, column=c, value=v)
    for j in range(len(meses)):
        col = L(f0 + j)
        wcab.cell(row=row, column=f0 + j, value=f"=SUM({col}{start}:{col}{row - 1})")
    paint(wcab, row, len(hs), R.TOTAL_FILL)
    resumen_cab.append((ceco, rend, float(cg[cg["Mes"].isin(meses)]["Precio Fnal"].sum()),
                        float(g[g["Concepto"] == "APPS_TRANSPORTE"]["Monto"].sum()),
                        float(g[g["Concepto"] == "COMBUSTIBLE"]["Monto"].sum())))
    row += 2
tT = f0 + len(meses)
for x in range(2, row):
    if wcab.cell(row=x, column=1).value:
        wcab.cell(row=x, column=tT, value=f"=SUM({L(f0)}{x}:{L(tT - 1)}{x})")
        wcab.cell(row=x, column=tT + 1, value=f"=IFERROR(AVERAGEIF({L(f0)}{x}:{L(tT - 1)}{x},\">0\"),\"\")")
        for c in range(f0, tT + 2):
            wcab.cell(row=x, column=c).number_format = MONEY
wcab.cell(row=row + 1, column=1, value=(
    "Usuarios de Cabify corporativo (Cabify Extracto, viajes no 'No encontrado', mismos meses que las rendiciones) que además "
    "rinden APPS_TRANSPORTE o COMBUSTIBLE según lo leído en los adjuntos. Match: hoja 'Auditoria Cabify vs Rend' o "
    "primer nombre + un apellido coincidentes (validar homónimos).")).font = Font(italic=True, size=9)
R._widths(wcab, [26, 30, 30, 10, 34] + [12] * len(meses) + [14, 13])

# ---------- 6. Base dinámica ----------
wbd = wb.create_sheet("Base Dinámica")
header(wbd, list(df.columns))
for i, rr in enumerate(df.itertuples(index=False), 2):
    for c, v in enumerate(rr, 1):
        wbd.cell(row=i, column=c, value=v)
    wbd.cell(row=i, column=list(df.columns).index("Monto") + 1).number_format = MONEY
wbd.auto_filter.ref = f"A1:{L(len(df.columns))}{len(df) + 1}"
R._widths(wbd, [28, 32, 9, 9, 22, 34, 44, 26, 30, 60, 40, 14, 11])

# ---------- 7. Desglose OTROS ----------
ot = df[df["Concepto"] == "OTROS"]
wo = wb.create_sheet("Desglose OTROS", 4)
row = 1
wo.cell(row=row, column=1, value="Qué hay dentro de OTROS (monto prorrateado a lo declarado)").font = Font(bold=True, size=12)
row += 2


def tabla(titulo, piv, row, first_w=None):
    wo.cell(row=row, column=1, value=titulo).font = Font(bold=True)
    row += 1
    header(wo, [piv.index.name or ""] + [str(c) for c in piv.columns], row)
    for i, (k, vals) in enumerate(piv.iterrows(), row + 1):
        wo.cell(row=i, column=1, value=str(k))
        for j, v in enumerate(vals, 2):
            c = wo.cell(row=i, column=j, value=float(v) if v == v else None)
            c.number_format = "0.0%" if "%" in str(piv.columns[j - 2]) else MONEY
    return row + len(piv) + 3


sub = ot.groupby("Subtipo OTROS")["Monto"].agg(["sum", "count"]).sort_values("sum", ascending=False)
sub.columns = ["Monto", "N° ítems"]
sub["% de OTROS"] = sub["Monto"] / sub["Monto"].sum()
sub.index.name = "Subtipo"
row = tabla("1. OTROS por subtipo", sub, row)
pc = ot.pivot_table(index="Subtipo OTROS", columns="CECO", values="Monto", aggfunc="sum").fillna(0)
pc = pc.loc[sub.index]
pc["Total"] = pc.sum(axis=1)
pc.index.name = "Subtipo \\ CECO"
row = tabla("2. Subtipo x CECO", pc, row)
pm = ot.pivot_table(index="Subtipo OTROS", columns="Mes", values="Monto", aggfunc="sum").fillna(0).loc[sub.index]
pm["Total"] = pm.sum(axis=1)
pm.index.name = "Subtipo \\ Mes"
row = tabla("3. Subtipo x Mes (evolutivo)", pm, row)
top_r = ot.groupby(["CECO", "Rendidor"])["Monto"].sum().sort_values(ascending=False).head(25)
pr = ot[ot.set_index(["CECO", "Rendidor"]).index.isin(top_r.index)].pivot_table(
    index=["Rendidor"], columns="Subtipo OTROS", values="Monto", aggfunc="sum").fillna(0)
pr["Total"] = pr.sum(axis=1)
pr = pr.sort_values("Total", ascending=False)
pr.index.name = "Rendidor (top 25 en OTROS)"
row = tabla("4. Rendidores con más OTROS x subtipo", pr, row)
wo.cell(row=row, column=1, value="5. Principales emisores / detalles dentro de cada subtipo (top 8)").font = Font(bold=True)
row += 1
header(wo, ["Subtipo", "Emisor / detalle", "Monto", "N° ítems"], row)
row += 1
for st in sub.index:
    g = ot[ot["Subtipo OTROS"] == st].groupby("Emisor / detalle")["Monto"].agg(["sum", "count"]).sort_values(
        "sum", ascending=False).head(8)
    for k, (v, n) in g.iterrows():
        for c, val in enumerate([st, k, float(v), int(n)], 1):
            wo.cell(row=row, column=c, value=val)
        wo.cell(row=row, column=3).number_format = MONEY
        row += 1
R._widths(wo, [60, 45] + [14] * 20)

# ---------- 8. Pendientes de detalle (para gestión/control) ----------
pend = df[((df["Concepto"] == R.SIN_ID) | df["Subtipo OTROS"].isin(GENERICOS))]
pend = pend.groupby(["ID", "CECO", "Rendidor"]).agg(Monto=("Monto", "sum"),
                                                     Probable=("Concepto probable (sin detalle)", "first"),
                                                     Base=("Base del probable", "first"),
                                                     Detalle=("Detalle de la rendición (carpetas/correos)", "first"),
                                                     Motivo=("Concepto detallado", "first")).reset_index()
pend = pend[pend["Monto"] > 0].sort_values("Monto", ascending=False)
link = {res["rendicion"]["id"]: res["rendicion"].get("adjunto") for res in results}
wpd = wb.create_sheet("Pendientes de detalle", 5)
hs = ["ID", "CECO", "Rendidor", "Monto sin detalle", "Concepto probable", "Base del probable", "Motivo / tipo",
      "Detalle de la rendición (carpetas/correos)", "Adjunto"]
header(wpd, hs)
for i, rr in enumerate(pend.itertuples(index=False), 2):
    vals = [rr.ID, rr.CECO, rr.Rendidor, rr.Monto, rr.Probable or "", rr.Base or "", rr.Motivo, rr.Detalle, "Abrir"]
    for c, v in enumerate(vals, 1):
        wpd.cell(row=i, column=c, value=v)
    wpd.cell(row=i, column=4).number_format = MONEY
    if link.get(rr.ID):
        wpd.cell(row=i, column=9).hyperlink = link[rr.ID]
        wpd.cell(row=i, column=9).font = Font(color="0563C1", underline="single")
wpd.auto_filter.ref = f"A1:I{len(pend) + 1}"
wpd.freeze_panes = "D2"
R._widths(wpd, [11, 26, 30, 14, 26, 30, 50, 80, 8])

wb.save(out)
json.dump({"archivo": str(out), "top_ceco": {k: [v[0], list(v[1].items())] for k, v in top_ceco.items()},
           "cabify": resumen_cab}, open("resumen.json", "w"), ensure_ascii=False, indent=1, default=float)
print("OK ->", out)
